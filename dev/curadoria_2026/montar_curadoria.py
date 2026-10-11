from pathlib import Path
import json, re, unicodedata, hashlib, csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment
ROOT = Path('D:/Claude/EAPA-Ecossistema/EAPADados')
OUT = ROOT/'dev/curadoria_2026'
if (OUT/'curadoria_eapadados_revisao.xlsx').exists():
    raise SystemExit('Planilha já existe: não sobrescrever decisões do autor. Gere nova versão com outro nome, se necessário.')
inv = json.loads((OUT/'inventario_tecnico.json').read_text(encoding='utf-8'))
P = {}
def put(n, ana, status, ev, priority='Média', nature='Real', action='Manter', restr=''):
    P[n] = dict(analise=ana,situacao=status,avaliacao=ev,prioridade=priority,natureza=nature,acao=action,restricao=restr)
put('abalone_adultos','Descritiva; comparação por sexo; potência; PCA','já ajustado','Manter 2.641 linhas e 9 colunas: recorte de adultos e conversão das medidas já documentados. Femea/Macho → fêmea/macho. Preservar anéis e componentes do peso; sem novo recorte por tamanho.')
put('aquicultura_br','Coroplético; comparação entre anos','já ajustado','Manter 216 linhas (27 UFs × 8 anos) e 5 colunas. posicao é derivada, mas útil ao ranking. Padronizar caixa dos estados, mantendo siglas oficiais. Oito anos por UF não sustentam decomposição sazonal mensal.')
put('artemia','Teste t independente; comparação de rações','precisa de ajuste','Propor retirada: a documentação declara dados fictícios. Substituir por exemplo real aprovado, como chickwts. Não apagar a linha repetida por padrão. A/B não identificam composição das rações.','Alta','Sintético — extra não autorizado','Retirar após aprovação')
put('bagley_lwr_central_america','Potência peso-comprimento; exploração por espécie','precisa de ajuste','Propor retirar log_peso_g e log_sl_mm (recalculáveis), nadadeira_dorsal (292/306 NA), nadadeira_caudal (298/306 NA) e sexo (271/306 NA). Manter coordenadas, localidade, país, espécie e no → id_registro. Preservar 306 linhas: ausência de comprimento_total_mm não invalida comprimento_padrao_mm.','Alta',restr='Não agrupar espécies numa curva única. Campos retirados continuam rastreáveis no bruto.')
put('biometria_caranguejos','Teste t; ANOVA 2 fatores; descritiva','precisa de ajuste','Manter 993 linhas e 5 colunas. Local/Sexo/Estacao → local/sexo/estacao; LC/CC → largura_carapaca/comprimento_carapaca, unidade pendente. Sexo → fêmea/macho; estação → chuvosa/seca. Não excluir as 551 linhas repetidas sem identificação da coleta.','Alta',restr='Confirmar unidade de LC/CC e protocolo/independência na fonte.')
put('brine_carbonatos','PCA; HCA; padronização','já ajustado','Manter 19 amostras e 8 colunas; correções V18/V19 já documentadas. Íons → bicarbonato_ppm, sulfato_ppm, cloreto_ppm, calcio_ppm, magnesio_ppm, sodio_ppm. Trocar underscores por espaços nos nomes próprios das formações geológicas.')
put('cabritos_fa_coltro','PCA; HCA; comparação de dietas','precisa de ajuste','Manter 32 animais e 10 colunas. n3/n6/sfa/mufa/pufa → omega_3_pct/omega_6_pct/saturados_pct/monoinsaturados_pct/poli_insaturados_pct; razões com prefixo razao_. T1…T4 → dieta_1…dieta_4, mantendo energia_mj_kg. Não inserir automaticamente razões junto com componentes na PCA.')
put('camarao_vannamei_biometria','Descritiva; peso-comprimento; comparação de origens','precisa de ajuste','Propor retirar arquivo_origem (guardar correspondência no dicionário); preservar ID, base, origem, amostra e imagem_completa para rastrear agrupamentos. Manter 170 linhas, inclusive 22 sem cefalotorax_cm. DB1…DB5 → fazenda_1, fazenda_2_coleta_1, fazenda_2_coleta_2, viveiro_coleta_1, viveiro_coleta_2, conforme arquivos de origem.')
put('camaroes_sexo','Qui-quadrado; uma e duas proporções','precisa de ajuste','Manter 170 indivíduos e 2 colunas. Sexo → fêmea/macho. Expandir P.brasiliensis/P.schimitti só após confirmar identificação e grafia na fonte; não inferir gênero pela inicial. Completar referência do livro. Não excluir as 166 linhas repetidas.','Alta',restr='Origem real declarada; referência bibliográfica ainda incompleta.')
put('cangulo_crescimento','Potência; transformação logarítmica','precisa de ajuste','Propor retirar apenas ln_peso e ln_comp. p1/p2/p3 → peso_1_g/peso_2_g/peso_3_g; peso_g → peso_medio_g. Manter 16 classes de comprimento; não empilhar os três pesos como medidas repetidas do mesmo peixe. Completar fonte/protocolo antes de chamar de exemplo real confirmado.','Alta','Real declarado — fonte incompleta',restr='Documentação diz que p1/p2/p3 são indivíduos por comprimento; conferir no original.')
put('captura_pescada_amarela','Preparo; captura por viagem; comparação sazonal','precisa de ajuste','Propor retirar cpue_pescada (252/252 NA) e cpue (251/252 NA). Manter dias_ao_mar (13 valores) e demais colunas até conferir a fonte. Manter 252 viagens. Propor data_saida/data_chegada e nomes tidy; unidades de precipitacao/peso_balanca/CPUE exigem confirmação.','Alta',restr='Não imputar CPUE nem excluir 239 viagens por ausência de dias_ao_mar. Não concatenar barcos como uma série.')
put('captura_pescada_amarela_bruta','Treino de limpeza com dados reais','precisa de ajuste','Preservar 252 linhas e 16 colunas se aprovada como exercício. Não renomear cabeçalhos nem corrigir categorias aqui: são alvos do preparo. Identificar defeitos de digitação inseridos de propósito. Alternativa do autor: retirar cópia bruta do pacote e mantê-la só como arquivo de exercício.','Média','Real com defeitos didáticos',restr='Não é base sintética, mas a cópia foi adulterada para treino; distinguir defeitos das medições reais.')
put('captura_petrechos','Contingência; contagem por petrecho','precisa de ajuste','Propor retirada: gerado por sample/rpois em preparar_dados.R. Usar dados reais aprovados para contingência/contagem; InsectSprays e ships são alternativas metodológicas.','Alta','Sintético — extra não autorizado','Retirar após aprovação')
put('cpue_tubarao','Série temporal; suavização; autocorrelação','precisa de ajuste','Manter 424 linhas e 5 colunas. Vessel/Year/Month/CPUE/Data → embarcacao/ano/mes/cpue/data. Conservar MB/MC/SB como IDs, sem inventar nomes de barcos. Completar referência/unidade de CPUE e conferir regularidade por barco antes de decompor ou prever.','Alta','Real declarado — fonte incompleta')
put('crabs_morfometria','Amostragem estratificada; ANCOVA; PCA','já ajustado','Manter 200 indivíduos e 10 colunas. Preservar IDs, estratos e classe_largura_carapaca; fora da PCA. Classes Q1…Q4 → quartil_1…quartil_4. Sexo → fêmea/macho. ANCOVA candidata exige interação e diagnóstico. Não reduzir para caber em tela.')
put('darter_ontario','Mann-Whitney; descritiva por rio','já ajustado','Manter 54 linhas e 3 colunas; nomes/unidades já claros. Manter Salmon/Trent como nomes próprios. Não excluir as 4 linhas repetidas. Registrar reconstrução de figura, sem afirmar que são medições brutas individuais.','Média','Real reconstruído de figura')
put('dicionario_variaveis_amostragem','Metadados das bases de amostragem','já ajustado','Manter 28 linhas e 6 colunas. Atualizar referências a variáveis/categorias só quando mudanças nas bases forem aprovadas. NA em unidade é esperado para IDs/categorias. Não é base de inferência.','Baixa','Tabela auxiliar')
put('doubs_ambiente','PCA; HCA; gradiente ambiental','precisa de ajuste','Manter 30 estações e 13 variáveis. Propor nomes completos com unidades (dfs → distancia_nascente_km, alt → altitude_m etc., aba Variáveis). Transformar rownames E01…E30 em coluna id_estacao. Manter trechos, inclusive duas estações impactadas; não excluir porque não permitem elipse.','Alta',restr='trecho é classificação didática derivada, não fator experimental; possível dependência ao longo do rio.')
put('estacoes_ictiofauna','Mapas de pontos e bolhas','precisa de ajuste','Propor retirada: coordenadas/CPUE/abundância sintéticas. Substituto real depende de base georreferenciada licenciada e protocolo; não converter grade NOAA em estações biológicas. Preservar intacto até aprovação.','Alta','Sintético — extra não autorizado','Retirar após aprovação')
put('exportacao_pargo','Arrumar largo → longo; comércio; série anual','precisa de ajuste','Manter 2005–2025 e 4 códigos NCM. Propor retirar Código ISIC Grupo e Descrição ISIC Grupo (constantes), guardando em metadados. Preservar país/UF/NCM/descrição. Formato longo proposto: pais, uf, codigo_ncm, descricao_ncm, ano, valor_fob_usd, massa_liquida_kg. Futuras 84 linhas = 4 códigos × 21 anos, sem descarte.','Média',restr='Confirmar vigência dos NCMs e se 2025 está completo; não somar códigos como série homogênea sem conferir.')
put('gammarus_dieta_temperatura_resumo','Exploração de médias; leitura de interação','já ajustado','Manter 12 combinações e 11 colunas de médias, DP e n. Dietas → fucus/folhas de cenoura/polpa de coco. Não reconstruir réplicas nem rodar ANOVA/Friedman nas médias como se fossem indivíduos.','Alta','Real — estatísticas agregadas',restr='Não cobre lacuna de ANOVA fatorial com observações individuais.')
put('idades_savel_repetibilidade','Wilcoxon pareado; repetibilidade','precisa de ajuste','Manter 53 peixes e 8 colunas. fish_id/true_age/reader_a_1…reader_c_2 → id_peixe/idade_real_anos/leitor_a_1_anos…leitor_c_2_anos. Preservar leitores; parear só leituras escolhidas na análise. Não excluir globalmente 20 peixes por ausência no leitor B.')
put('isoproteica_bagre','ANOVA 1 fator; Welch; Kruskal-Wallis','já ajustado','Manter 19 linhas e 2 colunas, já enxutas. Preservar ração como grupo; ver níveis na aba Categorias. Fonte Bhujel, Tabela 7.4 documentada; confirmar unidade experimental no original. Não balancear removendo observações.')
put('lagostas_kelp_sexo','Qui-quadrado; proporções por sítio','já ajustado','Manter 1.457 indivíduos e 8 colunas: filtro anterior por sexo/sítio registrado no script. ano/mes/mes_nome são redundantes mas úteis a exercícios/rótulos; manter. Códigos de sítio permanecem até obter nomes oficiais. Sexo → fêmea/macho. Registrar contagens do filtro anterior; sem novo corte.')
put('morfometria_barbo','Regressão linear simples; HCA','já ajustado','Manter 100 peixes: recorte balanceado já documentado. Propor retirar rio, rótulo didático 1:1 com populacao, guardando correspondência nos metadados. POP06…POP10 → população 06…10. Não acrescentar _mm às medidas corrigidas sem conferir escala alométrica.','Média',restr='Não selecionar novamente indivíduos para melhorar p-valores; diagnóstico documentado vale para este recorte.')
put('ocorrencias_peixes','Densidade espacial (módulo futuro)','precisa de ajuste','Propor retirada: coordenadas sintéticas. Densidade está em preparação no menu; substituto não é prioridade imediata. Futuras ocorrências reais devem documentar esforço/viés de coleta.','Alta','Sintético — extra não autorizado','Retirar após aprovação')
put('peixes_morfometria_multivariada','PCA; HCA; seleção de variáveis','precisa de ajuste','Propor retirar especie_codigo (repete especie). Manter 299 indivíduos, população, ID e 35 medidas. Não cortar medidas só por correlação: altera distâncias/componentes. Expandir códigos anatômicos só com dicionário original; sl/hmax/hmin têm tradução documentada. Demais códigos pendentes.','Alta',restr='Unidade mm documentada; significado de 32 códigos não resolvido na documentação local.')
put('peixes_nutricao_portugal','PCA; HCA; perfis nutricionais','já ajustado','Manter 40 linhas e 22 colunas. Traduzir sfa/mufa/pufa/epa/dha com unidades. cluster_autores e perfil_nutricional_autores são rótulos externos, fora da entrada PCA/HCA para evitar circularidade. Não excluir atípicos automaticamente; padronizar rótulo atípicos com acento.')
put('pinguins','Descritiva; teste t; ANOVA; ANCOVA; PCA','já ajustado','Manter 344 aves e 8 colunas; nomes já traduzidos. Conservar 11 ausências de sexo e 2 de biometria: casos completos só no ramo que exigir essas colunas. Manter nomes próprios de espécies/ilhas; sexo → fêmea/macho. Avaliar grupos antes de ajustar reta única.')
put('qfasa_acidos_graxos','Metadados; seleção de ácidos graxos','precisa de ajuste','Objeto real é data.frame 39 × 1 com coluna FA; documentação diz vetor. Corrigir documentação preservando formato. FA → acido_graxo; revisar correspondência com as duas bases sem traduzir nomenclatura química.','Alta','Tabela auxiliar')
put('qfasa_predadores','PCA/HCA de assinaturas lipídicas (avançado)','precisa de ajuste','Manter 10 amostras e 70 colunas até validar painel; são apenas 4 animais com biópsias repetidas. SampleCode/AnimalCode/SampleGroup/Biopsy → id_amostra/id_animal/grupo_amostra/biopsia. Normalizar ácidos só com de-para bijetivo e sincronizado às presas/lista FA. Não cortar só porque p > n.','Alta',restr='Não tratar 10 biópsias como 10 animais independentes; nomes/unidades químicos pendentes.')
put('qfasa_presas','PCA/HCA de assinaturas lipídicas (avançado)','precisa de ajuste','Manter 302 amostras e 70 colunas. Lab.Code/Species/lipid → id_amostra/especie/lipideos (unidade pendente). Não igualar c7Me16.0/c7Mec16.0 ou c16.3w1/c16.3w4 sem conferir fonte/lista FA. Traduzir nomes comuns só com espécie confirmada.','Alta',restr='Painéis de presas/predadores diferem; documentar unidade das assinaturas e de lipid.')
put('recifes_ostras_heatmap','Correlação; heatmap; PCA; contagens','precisa de ajuste','Propor retirar avaliador_video, mantendo metadados. Preservar IDs de recife/levantamento, sítio, zona, estação e duração de vídeo; manter 301 linhas e respostas ecológicas/geométricas. Centre/Edge/Off → centro/borda/fora do recife; Summer/Winter → verão/inverno. Confirmar unidades de áreas/perímetro/distâncias.','Alta',restr='Repetição por recife/levantamento; duração/visibilidade podem afetar contagens. Não afirmar independência.')
put('referencias_dados_amostragem','Referências das bases de amostragem','já ajustado','Manter 5 linhas e 4 colunas. Preservar títulos bibliográficos e nomes de pacotes; atualizar apenas referências cruzadas aprovadas. Tabela auxiliar sem análise estatística.','Baixa','Tabela auxiliar')
put('salvelino_formalina_remocao','ANOVA fatorial 2 × 2; sobrevivência','já ajustado','Manter 30 unidades e 7 colunas: seleção do registro final já documentada. formalin → formalina; sem/com → sem formalina/com formalina; remoção → sem remoção/com remoção. Preservar contagens para conferir porcentagens. Não reduzir unidades para balancear.','Alta',restr='Denominadores diferentes exigem diagnóstico; ANOVA de porcentagem não é garantia. Recorte final não serve a medidas repetidas.')
put('snapper_hg','Amostragem; von Bertalanffy; idade-comprimento','já ajustado','Manter 1.331 linhas e 6 colunas. Filtro anterior reteve idade atribuída; não reinserir registros reconstruídos sem curadoria. tem_idade_atribuida é constante mas documenta filtro; manter. Classes Q1…Q4 → quartil_1…quartil_4. Sem novo corte de indivíduos.')
put('sst_costa_norte','Raster ambiental (módulo futuro)','precisa de ajuste','Propor retirada: grade sintética. NOAA ERSST é candidato real reconstruído para TSM/séries, sem clorofila ou profundidade. Não substituir variáveis faltantes por valores fabricados.','Alta','Sintético — extra não autorizado','Retirar após aprovação')
put('tilapia_crescimento','Crescimento; linhas por tratamento','precisa de ajuste','Propor retirada: gerado por rnorm em preparar_dados.R. Sem ID de tanque/peixe; não serve para medidas repetidas reais. Orthodont ou dados de peixes aprovados cobrem melhor a estrutura temporal.','Alta','Sintético — extra não autorizado','Retirar após aprovação')
put('tilapia_microalgas','ANOVA de tratamentos; PCA exploratória','precisa de ajuste','Propor retirar tratamento_codigo (repete ID de tanque) e repeticao, guardando correspondência em dicionário. Preservar id_tanque e respostas. 15 tanques: não ajustar modelo múltiplo extenso da documentação. Expandir sgr/tan/srp/rbc/pcv/wbc/bun com unidades.','Alta',restr='Controle não cruza com 25/50: não é fatorial completo tipo × dose. PCA com 15 linhas é exploratória; múltipla ainda em desenvolvimento na tela.')
put('treino_coletas','Separar colunas; preparo','pronto','Manter 15 linhas e 5 colunas, inclusive duas compostas; separá-las antes destruiria exercício. Aviso obrigatório: DADOS SINTÉTICOS DE TREINO — não usar para inferência científica. Exceção autorizada, junto com treino_desembarque.','Alta','Sintético — treino autorizado')
put('treino_desembarque','Arrumar largo → longo','pronto','Manter 12 linhas e 8 colunas largas com cabeçalhos compostos. Nomes tidy só na saída: porto, especie, ano, captura_t, receita_mil_reais (confirmar moeda fictícia do exercício). Aviso obrigatório: DADOS SINTÉTICOS DE TREINO.','Alta','Sintético — treino autorizado')
put('truta_riacho_crescimento','ANCOVA candidata; ração e temperatura','já ajustado','Manter 69 peixes e 8 colunas; preservar tanque. racao_fator high/low → alta/baixa. Manter massa inicial e temperatura. Não usar 69 peixes como réplicas independentes de temperatura/ração sem conferir alocação por tanque.','Alta',restr='ANCOVA simples da tela pode não representar efeito de tanque; conferir fonte antes de adotar exemplo canônico.')
put('walleye_erie','Amostragem simples/estratificada; crescimento','já ajustado','Manter 33.734 indivíduos e 12 colunas: útil como população de amostragem. Não cortar por tamanho de tela. Preservar lance/grade/ano/região e 413 ausências de peso. Classes → quartil_1…quartil_4; sexo → fêmea/macho. Documentar limites de classes.','Alta',restr='Indivíduos do mesmo lance podem ser dependentes para inferência; amostragem didática não apaga a hierarquia.')
assert set(P)==set(inv),(set(inv)-set(P),set(P)-set(inv))
docs={}
for f in (ROOT/'R').glob('*.R'):
    txt=f.read_text(encoding='utf-8-sig',errors='replace')
    for n in inv:
        if ('"'+n+'"' in txt or '@name '+n in txt) and (f.name.startswith('data') or f.name in ['qfasa_dados.R','camarao_vannamei_biometria.R']):
            docs.setdefault(n,[]).append(f.relative_to(ROOT).as_posix())
rows=[]
for i,(n,x) in enumerate(inv.items(),1):
    rows.append(dict(nome=n,origem='EAPADados',**P[n],id=f'E{i:03}',dimensao=f"{x['linhas']:,} × {x['colunas']}".replace(',','.'),decisao='Pendente',autor='',fonte_local='; '.join(docs.get(n,[]))))
new=[]
def candidate(n,orig,ana,ev,url,priority='Alta',nature='Real',restr='',size='A conferir',status='precisa de ajuste'):
    new.append(dict(nome=n,origem=orig,analise=ana,situacao=status,avaliacao=ev,prioridade=priority,natureza=nature,acao='Incorporar após aprovação',restricao=restr,fonte_local='',url=url,dimensao=size,decisao='Pendente',autor=''))
candidate('whitefish_salinidade','Artigo / Dryad','Subamostras por tanque; crescimento; salinidade','Candidato com peixes reais para subamostras. Examinar File_D.csv/README e selecionar, após aprovação, um dia com tanques replicados por salinidade. Preservar tanque de origem e atual. Nomes propostos: salinidade_ppt, comprimento_mm, peso_g, dia, id_tanque. Arquivo ainda não baixado/validado numericamente.','https://doi.org/10.5061/dryad.bg79cnpph',restr='Peixes não marcados individualmente: não é pareamento individual longitudinal. Conferir replicação independente, licença e confundimento temperatura/tempo.')
candidate('mesocosmos_secagem','Artigo / Dryad','ANOVA fatorial; longitudinal por mesocosmo','Experimento com 64 mesocosmos, 4 tratamentos tróficos × 2 regimes de secagem. Para peixes: 32 tanques, um peixe por tanque. Inspecionar Fish_Data_Dryad.xlsx e scripts antes de selecionar respostas/nomes. Não propor Gompertz como funcionalidade pronta. Arquivo ainda não baixado.','https://doi.org/10.5061/dryad.hx3ffbgkn',restr='Preservar tanque/tempo para registros repetidos; confirmar licença e unidades antes de redistribuir.')
candidate('tsm_noaa_costa_norte','Base aberta / NOAA','Série mensal; ACF; decomposição','Propor recorte fixo 2000–2023 de ERSSTv6, regra espacial a aprovar: célula oceânica ou média de área da costa norte. Saída data, latitude_graus, longitude_graus, temperatura_c. Registrar versão/procedimento; sem download global ou atualização automática.','https://www.ncei.noaa.gov/products/extended-reconstructed-sst',priority='Média',nature='Real — produto reconstruído',restr='Reconstrução estatística a partir de observações, não medição direta de estação; não fornece clorofila/profundidade. Raster ainda em preparação.')
baseurl='https://stat.ethz.ch/R-manual/R-devel/library/'
rmeta=json.loads((OUT/'candidatos_r_verificados.json').read_text(encoding='utf-8'))
def rcan(n,p,ana,ev,restr='',priority='Alta'):
    m=rmeta.get(p+'::'+n)
    size=f"{m['linhas']} × {m['colunas']}" if m else '144 × 4 (manual)'
    url=baseurl+p+'/html/'+n+'.html' if p!='agridat' else 'https://search.r-project.org/CRAN/refmans/agridat/html/grover.rcb.subsample.html'
    candidate(p+'::'+n,'Pacote R: '+p,ana,ev,url,priority,restr=restr,size=size,status='pronto')
rcan('Orthodont','nlme','Medidas repetidas; Friedman','27 indivíduos × 4 idades (8, 10, 12, 14), grade completa. distance/age/Subject/Sex → distancia_mm/idade_anos/id_individuo/sexo. Male/Female → masculino/feminino. Na tela: idade como momento e ID como sujeito. Sexo é estratificação opcional.','Pronto na origem; tradução depende de aprovação. Tempo ordenado não é tratamento randomizado. Diagnosticar resíduos/esfericidade.')
rcan('bacteria','MASS','McNemar (pares binários)','Selecionar semanas 0 e 11, alargar por ID. Verificados 44 pares de 50 crianças: 2 não/não, 2 não/sim, 10 sim/não, 30 sim/sim. ID/y/week → id_individuo/presenca_bacteria/semana; n/y → ausente/presente. Saída: presenca_inicial/presenca_final.','Pronto na origem; recorte é proposta. Registrar seis perdas, sem imputação. Não tratar medidas repetidas como independentes.')
rcan('ships','MASS','Poisson; exposição/offset','40 células de frota. type/year/period/service/incidents → tipo_navio/periodo_construcao/periodo_operacao/servico_meses/incidentes. Filtrar seis células sem exposição apenas no ramo com offset. Construção 60/65/70/75 → 1960–1964/1965–1969/1970–1974/1975–1979; operação 60/75 → 1960–1974/1975–1979. Poisson: dispersão ≈ 1,69.','Contagens agregadas, não 40 navios. NB com tipo e períodos fatoriais atingiu limite de iterações de theta; não adotar como exemplo NB garantido. Usar InsectSprays como alternativa verificada.')
rcan('kyphosis','rpart','Logística binária; uma proporção','81 pacientes. Kyphosis/Age/Number/Start → cifose/idade_meses/numero_vertebras/vertebra_inicial. absent/present → ausente/presente; evento presente=1 no ramo logístico. Manter todos os registros.','Uso metodológico; vertebra_inicial é posição anatômica. Avaliar separação/incerteza; não afirmar desempenho clínico.')
rcan('ToothGrowth','datasets','ANOVA fatorial 2 × 3; doses','60 animais, dois suplementos × três doses. len/supp/dose → comprimento_odontoblasto/suplemento/dose_mg_dia; OJ/VC → suco de laranja/ácido ascórbico. dose como fator para ANOVA; contínua só para tendência.','Unidade de len ausente na ajuda: não inventar _mm. Confirmar artigo antes de incorporação definitiva; dados originais utilizáveis.')
rcan('sleep','datasets','Teste t pareado; Wilcoxon pareado','20 registros de dez pessoas sob dois fármacos. extra/group/ID → aumento_sono_h/farmaco/id_individuo. 1/2 → fármaco 1/fármaco 2. Alargar por ID se a tela pedir duas respostas.','Condições nas mesmas pessoas, não grupos independentes. Não parear só pela ordem das linhas.',priority='Média')
rcan('InsectSprays','datasets','Poisson; NB; Kruskal-Wallis','72 unidades e seis inseticidas. count/spray → contagem_insetos/inseticida; A…F → inseticida A…F, sem inventar composição. No motor da TRILHA: Poisson com dispersão ≈ 1,51; NB convergiu com theta ≈ 28,10. Manter observações.','Contagem não garante ANOVA normal. Exposição ausente na tabela: não inventar offset. Conferência de ajuste não dispensa diagnóstico.',priority='Alta')
rcan('Orange','datasets','Curva logística; exploração longitudinal','35 registros: cinco árvores × sete tempos. Tree/age/circumference → id_arvore/idade_dias/circunferencia_mm. Preservar base; ajuste introdutório pode mostrar árvore 3, como a ajuda, sem apagar demais árvores do pacote.','Mesma árvore medida no tempo; tela não oferece modelo não linear misto.',priority='Média')
rcan('co2','datasets','Decomposição; ACF; série mensal','468 meses (1959–1997), objeto ts. Converter índice para data e medida para co2_ppm. Registrar três meses interpolados na origem (fev–abr/1964); não excluir meses e quebrar regularidade.','Série histórica congelada, não medição atual. Interpolação publicada não autoriza fabricar dados.',priority='Média')
rcan('chickwts','datasets','ANOVA; Welch; Kruskal-Wallis','71 aves. weight/feed → peso_g/racao. horsebean/linseed/soybean/sunflower/meatmeal/casein → fava/linhaça/soja/girassol/farinha de carne/caseína. Alternativa real simples a artemia.','Não balancear grupos removendo aves; pressupostos sujeitos a diagnóstico.',priority='Média')
rcan('grover.rcb.subsample','agridat','Subamostras em blocos (limite da tela)','144 linhas: nove fertilizantes × quatro blocos × quatro subamostras. tiller/trt/block/unit → perfilhos/tratamento/bloco/subamostra. Criar id_parcela=bloco+tratamento. Rótulos completos de tratamentos dependem do livro.','Pronto na origem, conferido no manual, não instalado. Efeito de bloco e resposta de contagem: tela atual de um fator + unidade não representa todo delineamento. Complementar, não canônico imediato.',priority='Baixa')
rcan('quakes','datasets','Mapas de pontos; bolhas proporcionais','1.000 eventos sísmicos reais próximos de Fiji. lat/long/depth/mag/stations → latitude_graus/longitude_graus/profundidade_km/magnitude/estacoes_detectoras. Alternativa metodológica real para substituir estações sintéticas. Preservar todos os eventos e coordenadas originais.','Desativar fundo de estados brasileiros na tela. Longitudes cruzam 180°: manter escala leste contínua da fonte para este desenho sem fundo. Não inferir esforço biológico de estações detectoras. Não é base de pesca.',priority='Média')
for i,r in enumerate(new,1):r['id']=f'N{i:03}'
rows+=new
M={
 'bagley_lwr_central_america':{'no':'id_registro','log_sl_mm':'log_comprimento_padrao'},
 'biometria_caranguejos':{'Local':'local','Sexo':'sexo','Estacao':'estacao','LC':'largura_carapaca','CC':'comprimento_carapaca'},
 'brine_carbonatos':dict(zip(['HCO3','SO4','Cl','Ca','Mg','Na'],['bicarbonato_ppm','sulfato_ppm','cloreto_ppm','calcio_ppm','magnesio_ppm','sodio_ppm'])),
 'cabritos_fa_coltro':{'n3':'omega_3_pct','n6':'omega_6_pct','n6_n3':'razao_omega_6_3','sfa':'saturados_pct','mufa':'monoinsaturados_pct','pufa':'poli_insaturados_pct','pufa_sfa':'razao_poli_saturados'},
 'cangulo_crescimento':{'p1':'peso_1_g','p2':'peso_2_g','p3':'peso_3_g','peso_g':'peso_medio_g','ln_peso':'log_peso_medio','ln_comp':'log_comprimento'},
 'cpue_tubarao':{'Vessel':'embarcacao','Year':'ano','Month':'mes','CPUE':'cpue','Data':'data'},
 'captura_pescada_amarela':{'data_de_saida':'data_saida','data_de_chegada':'data_chegada'},
 'doubs_ambiente':dict(zip(['dfs','alt','slo','flo','pH','har','pho','nit','amm','oxy','bdo'],['distancia_nascente_km','altitude_m','declividade_por_mil','vazao_m3_s','ph','dureza_mg_l','fosfato_mg_l','nitrato_mg_l','amonio_mg_l','oxigenio_mg_l','dbo_mg_l'])),
 'idades_savel_repetibilidade':{'fish_id':'id_peixe','true_age':'idade_real_anos',**{f'reader_{a}_{b}':f'leitor_{a}_{b}_anos' for a in 'abc' for b in (1,2)}},
 'peixes_morfometria_multivariada':{'sl':'comprimento_padrao_mm','hmax':'altura_maxima_mm','hmin':'altura_minima_mm'},
 'peixes_nutricao_portugal':{'sfa_g_100g':'saturados_g_100g','mufa_g_100g':'monoinsaturados_g_100g','pufa_g_100g':'poli_insaturados_g_100g','epa_g_100g':'acido_eicosapentaenoico_g_100g','dha_g_100g':'acido_docosahexaenoico_g_100g'},
 'qfasa_predadores':{'SampleCode':'id_amostra','AnimalCode':'id_animal','SampleGroup':'grupo_amostra','Biopsy':'biopsia'},
 'qfasa_presas':{'Lab.Code':'id_amostra','Species':'especie','lipid':'lipideos'},
 'qfasa_acidos_graxos':{'FA':'acido_graxo'},'salvelino_formalina_remocao':{'formalin':'formalina'},
 'sst_costa_norte':{'lon':'longitude_graus','lat':'latitude_graus','sst':'temperatura_c','clorofila':'clorofila_mg_m3','profundidade':'profundidade_m'},
 'tilapia_crescimento':{'Tratamento':'tratamento','Semana':'semana','PesoMedio':'peso_medio','DataColeta':'data_coleta'},
 'tilapia_microalgas':{'sgr_pct_dia':'crescimento_especifico_pct_dia','tan_mg_l':'nitrogenio_amoniacal_total_mg_l','no2_n_mg_l':'nitrogenio_nitrito_mg_l','srp_mg_l':'fosforo_reativo_mg_l','rbc_10_6_ul':'eritrocitos_milhoes_ul','pcv_pct':'hematocrito_pct','wbc_10_3_ul':'leucocitos_milhares_ul','bun_mg_dl':'nitrogenio_ureico_mg_dl'}
}
def tidy(s):
    s=unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().lower()
    s=re.sub(r'[^a-z0-9]+','_',s).strip('_')
    return 'v_'+s if s and s[0].isdigit() else s
remove={'bagley_lwr_central_america':{'log_peso_g','log_sl_mm','nadadeira_dorsal','nadadeira_caudal','sexo'},'camarao_vannamei_biometria':{'arquivo_origem'},'cangulo_crescimento':{'ln_peso','ln_comp'},'captura_pescada_amarela':{'cpue_pescada','cpue'},'morfometria_barbo':{'rio'},'peixes_morfometria_multivariada':{'especie_codigo'},'recifes_ostras_heatmap':{'avaliador_video'},'tilapia_microalgas':{'tratamento_codigo','repeticao'}}
intentional={'treino_coletas','treino_desembarque','captura_pescada_amarela_bruta','exportacao_pargo'}
varrows=[];catrows=[]
for n,x in inv.items():
    for c in x['colunas_detalhe']:
        old=c['nome'];newname=M.get(n,{}).get(old,tidy(old));pending=False
        note='Manter; nome já tidy.' if old==newname else 'Padronizar nome; não converter valores.'
        if old in remove.get(n,set()):note='Retirada proposta na tabela principal; renomear só se o autor decidir manter.'
        if n in intentional:newname=old;note='Cabeçalho intencional: preservar no original. Nomes tidy na saída do preparo, conforme tabela principal.'
        if n=='peixes_morfometria_multivariada' and len(old)<=4 and old not in M[n]:
            newname='';note='PENDÊNCIA: recuperar significado anatômico do código; mm confirmado. Não inventar tradução.';pending=True
        if n=='biometria_caranguejos' and old in ['LC','CC']:note='PENDÊNCIA: confirmar unidade para acrescentar sufixo.';pending=True
        if n=='captura_pescada_amarela' and old in ['precipitacao','peso_balanca']:note='PENDÊNCIA: confirmar unidade na fonte.';pending=True
        if n=='cpue_tubarao' and old=='CPUE':note='PENDÊNCIA: confirmar denominador/unidade de CPUE.';pending=True
        if n=='qfasa_presas' and old=='lipid':note='PENDÊNCIA: confirmar unidade de lipid.';pending=True
        if n.startswith('qfasa_') and old not in M.get(n,{}):note='Grafia química normalizada; unidade/equivalência entre painéis pendentes. Revisão conjunta obrigatória.';pending=True
        if n=='morfometria_barbo' and old not in ['id_peixe','populacao','rio']:note='Medida corrigida alometricamente: não adicionar unidade por inferência.'
        if n=='recifes_ostras_heatmap' and old in ['area_2d','area_3d','perimetro','distancia_media_recifes_proximos','distancia_vizinho_mais_proximo','razao_perimetro_area','indice_para']:note='PENDÊNCIA: conferir unidade na fonte antes do sufixo.';pending=True
        varrows.append([n,old,newname,'Herdar do conjunto','',note,c['classe'],c['na'],c['distintos'],'Sim' if pending else 'Não'])
        lev=c.get('niveis');lev=lev if isinstance(lev,list) else [lev] if isinstance(lev,str) else []
        if 'factor' not in c['classe'] and not (c['classe']=='character' and lev and old in ['estado','uf','origem','grupo','racao','tipo_microalga','especie']):continue
        for level in lev:
            proposed=level;why='Manter nome próprio, rótulo claro ou identificador original.';needs=False;lo=level.lower()
            if lo in ['female','femea','fêmea']:proposed='fêmea';why='Padronizar sexo.'
            elif lo in ['male','macho']:proposed='macho';why='Padronizar sexo.'
            elif n=='brine_carbonatos':proposed=level.replace('_',' ');why='Nome próprio da formação, com espaços.'
            elif n=='cabritos_fa_coltro' and old=='tratamento':proposed='dieta_'+level[1:];why='Código legível; energia em coluna própria.'
            elif n=='salvelino_formalina_remocao':proposed=level+(' formalina' if old=='formalin' else ' remoção');why='Explicitar significado da categoria.'
            elif n=='recifes_ostras_heatmap' and old=='zona_recife':proposed={'Centre':'centro','Edge':'borda','Off':'fora do recife'}.get(level,level);why='Traduzir zona do recife.'
            elif n=='recifes_ostras_heatmap' and old=='estacao':proposed={'Summer':'verão','Winter':'inverno'}.get(level,level);why='Traduzir estação.'
            elif n=='truta_riacho_crescimento' and old=='racao_fator':proposed={'high':'alta','low':'baixa'}.get(level,level);why='Traduzir oferta de ração.'
            elif n=='gammarus_dieta_temperatura_resumo' and old=='dieta':proposed=level.lower().replace('_',' ');why='Padronizar dieta.'
            elif old.startswith('classe_') and re.search(r'_Q[1-4]$',level):proposed='quartil_'+level[-1];why='Explicitar classe; documentar limites numéricos no dicionário.'
            elif old=='populacao' and level.startswith('POP'):proposed='população '+level[3:];why='Explicitar código sem inventar localização.'
            elif n=='camarao_vannamei_biometria' and old=='base':proposed={'DB1':'fazenda_1','DB2':'fazenda_2_coleta_1','DB3':'fazenda_2_coleta_2','DB4':'viveiro_coleta_1','DB5':'viveiro_coleta_2'}[level];why='Conforme arquivo de origem documentado.'
            elif n=='biometria_caranguejos' and old=='Estacao':proposed=lo;why='Padronizar caixa.'
            elif n=='captura_pescada_amarela' and old in ['mes','nome_aparelho_pesca']:proposed=lo;why='Padronizar caixa; preservar ordem dos meses.'
            elif n=='camaroes_sexo' and old=='especie':proposed='';why='PENDÊNCIA: confirmar nome científico e grafia na fonte.';needs=True
            elif n=='qfasa_presas' and old=='Species':proposed='';why='PENDÊNCIA: nomes comuns ingleses ambíguos; confirmar espécie antes de traduzir.';needs=True
            elif n=='qfasa_predadores' and old=='SampleGroup':proposed='';why='PENDÊNCIA: significado de T/T1/T2 não documentado localmente.';needs=True
            elif n=='bagley_lwr_central_america' and old=='nadadeira_dorsal':proposed={'plain fin (no red)':'sem vermelho','red fin':'vermelha'}.get(level,level);why='Tradução; coluna com retirada proposta.'
            elif n=='bagley_lwr_central_america' and old=='nadadeira_caudal':proposed='nadadeira com ponto';why='Tradução; coluna com retirada proposta.'
            elif n=='peixes_nutricao_portugal' and old=='perfil_nutricional_autores':proposed={'atipicos':'atípicos','gordos':'gordos','magros':'magros'}.get(level,level);why='Rótulos externos; não entram como variáveis da PCA.'
            elif n=='aquicultura_br' and old=='estado':proposed=level.lower();why='Padronizar caixa dos estados; manter siglas oficiais em uf.'
            if n in intentional:proposed=level;why='Preservar categoria original do exercício; alterações só na saída do preparo.'
            catrows.append([n,old,level,proposed,'Herdar do conjunto','',why,'Sim' if needs else 'Não'])
wb=Workbook();guide=wb.active;guide.title='Comece aqui'
guide.append(['Curadoria EAPADados • 10/10/2026','Como decidir'])
instructions=[
('Escopo','43 objetos atuais + 15 candidatos = 58 entradas numa tabela. Nenhum conjunto foi alterado.'),
('1. Abra Curadoria','Cinco primeiras colunas seguem seu pedido. Filtre Prioridade=Alta para começar. Pacotes R estão ao fim, marcados pronto (na origem).'),
('2. Clique em Decisão','Pendente, Aprovar proposta, Aprovar com ajustes, Recusar ou Adiar. Todas começam Pendente: pronto não significa aprovado.'),
('3. Escreva sua orientação','Avaliação final do autor registra ajustes à proposta, que fica preservada. Aprovar com ajustes exige texto.'),
('4. Variáveis e Categorias','De-para de todas as colunas e níveis. Herdar do conjunto evita centenas de cliques: segue a decisão da linha principal. Você pode recusar/ajustar itens específicos. PENDÊNCIA técnica continua bloqueando o item mesmo com aprovação geral.'),
('5. Salve o XLSX','Ctrl+S no Excel ou OnlyOffice e avise que concluiu. Lerei o mesmo arquivo salvo e suas decisões. Sem macros ou aplicação automática de alterações.'),
('Exemplo ilustrativo','Decisão: Aprovar com ajustes | Avaliação final do autor: manter sexo; retirar só as duas colunas logarítmicas. Exemplo, não aprovação real.'),
('Interpretação das decisões','Recusar/Adiar/Pendente não autorizam ação. Aprovar proposta autoriza texto da linha. Aprovar com ajustes usa seu texto, conciliado com pendências/de-para antes de alteração.'),
('Limpeza × padronização','já ajustado registra curadoria anterior; pode haver nomes a padronizar. Não excluir linhas, balancear grupos ou imputar ausências sem decisão explícita.'),
('Sintéticos autorizados','Assumi treino_coletas e treino_desembarque como os dois citados. Seis extras identificados: artemia, captura_petrechos, estacoes_ictiofauna, ocorrencias_peixes, sst_costa_norte, tilapia_crescimento. Retirada só proposta.'),
('Tipos de dado','Sintético não se confunde com dado real agregado, reconstruído de figura ou produto ambiental reconstruído. Avisos específicos por conjunto.'),
('Espaço ocupado','Todos os .rda somam 337.401 bytes. Prioridade é clareza/utilidade, não reduzir população de amostragem para caber na tela.'),
('Validação','Objetos carregados isoladamente em R 4.6.1; dimensões/classes/níveis/NA/repetições conferidos. Onze candidatos em pacotes instalados; grover.rcb.subsample no manual. Literatura/NOAA ainda sem download.'),
('ClaRa hoje','comparar_medias e comparar_medianas disponíveis; outras perguntas planejadas. Propostas seguem métodos da TRILHA, sem prometer integração já feita.'),
('Fontes','Diagnóstico contém documentação local, URLs e restrições. Arquivos fonte_*.txt guardam ajuda dos pacotes instalados. Licença/versão devem acompanhar incorporação futura.'),
('Documentação obrigatória','Cada conjunto modificado deve declarar adaptação, identificar fonte original (pacote/objeto/versão ou artigo/base, DOI/URL) e resumir colunas removidas, filtros de linhas com n antes/depois, renomeações de variáveis/categorias e mudanças de formato/unidade. Registrar só alterações aplicadas; atualizar roxygen, man, dimensões e exemplos. Sem exclusão de linhas, declarar preservação.'),
('Aprovação e pendências','Aprovação geral não resolve unidade/nome desconhecido, licença incerta ou conflito entre abas. Casos continuam pendentes até definição explícita. Nomes não serão adivinhados.'),
('Arquivos técnicos','inventario_tecnico.json, candidatos_r_verificados.json, manifesto_sha256.json e propostas.json rastreiam levantamento. exportar_decisoes.py só gera relatório, sem mudar dados.')]
for r in instructions:guide.append(r)
headers=['nome do conjunto','origem','tipo de análise atendida','situação','avaliação resumida do que vai mudar','Decisão','Avaliação final do autor','ID','Prioridade','Dimensão atual','Natureza']
ws=wb.create_sheet('Curadoria');ws.append(headers)
for r in rows:ws.append([r['nome'],r['origem'],r['analise'],r['situacao'],r['avaliacao'],r['decisao'],r['autor'],r['id'],r['prioridade'],r['dimensao'],r['natureza']])
vs=wb.create_sheet('Variáveis');vs.append(['Conjunto','Nome atual','Nome proposto','Decisão','Ajuste do autor','Motivo / unidade / condição','Classe','NA','Valores distintos','Pendência técnica'])
for r in varrows:vs.append(r)
cs=wb.create_sheet('Categorias');cs.append(['Conjunto','Variável atual','Rótulo atual','Rótulo proposto','Decisão','Ajuste do autor','Motivo','Pendência técnica'])
for r in catrows:cs.append(r)
ds=wb.create_sheet('Diagnóstico');ds.append(['ID','Conjunto','Ação proposta','Restrição metodológica / fonte','Evidência e fonte local','Fonte externa','Linhas','Colunas','Linhas repetidas','Tamanho .rda (bytes)','Colunas ausentes ou constantes'])
for r in rows:
    x=inv.get(r['nome'],{})
    alerts='; '.join(f"{c['nome']}: {c['na']}/{x['linhas']} NA, {c['distintos']} valores" for c in x.get('colunas_detalhe',[]) if c['na'] or c['distintos']<=1)
    ds.append([r['id'],r['nome'],r['acao'],r['restricao'],r.get('fonte_local',''),r.get('url',''),x.get('linhas'),x.get('colunas'),x.get('duplicadas'),x.get('bytes'),alerts])
    cell=ds.cell(ds.max_row,6)
    if cell.value:cell.hyperlink=cell.value
gap=wb.create_sheet('Cobertura');gap.append(['Análise / funcionalidade','Cobertura atual','Diagnóstico','Proposta prioritária','Limite / condição','Evidência TRILHA / ClaRa'])
coverage=[
('Teste t; ANOVA 1 fator; Welch','isoproteica_bagre; pinguins; caranguejos','Coberto; substituir artemia fictícia','datasets::chickwts','Independência/variâncias exigem diagnóstico.','modules/mod_anova.R; ClaRa/R/clara_medias.R'),
('Mann-Whitney; Kruskal-Wallis','darter_ontario; isoproteica_bagre','Coberto','chickwts como alternativa','Darter reconstruído de figura.','modules/mod_nonparametric.R; ClaRa/R/clara_medianas.R'),
('Teste t / Wilcoxon pareados','idades_savel_repetibilidade','Coberto com ausências por leitor','datasets::sleep','Parear por ID.','modules/mod_parametric.R; mod_nonparametric.R'),
('Medidas repetidas','Sem exemplo canônico temporal completo','Lacuna alta','nlme::Orthodont','27 sujeitos × 4 idades completos; esfericidade/resíduos.','tests/test_parametricos_complementares.R; R/analises.R'),
('Friedman','Leituras de idade incompletas por leitor','Lacuna de grade completa','nlme::Orthodont','Mesmo sujeito em todos os momentos; não usar médias Gammarus.','tests/test_friedman_mcnemar.R'),
('McNemar','Sem dois estados binários pareados documentados','Lacuna alta','MASS::bacteria','44 pares semanas 0/11; seis perdas; alargar por ID.','modules/mod_pareados_categoricos.R; tests/test_friedman_mcnemar.R'),
('ANOVA com subamostras','truta tem tanque; alocação requer conferência','Lacuna alta de exemplo simples','whitefish_salinidade (a conferir)','Grover tem blocos/contagens, fora do modelo completo da tela atual.','R/anova_mista.R; modules/mod_anova_mista.R'),
('ANOVA 2 fatores','salvelino_formalina_remocao','Coberto; ampliar além de porcentagens','datasets::ToothGrowth','Gammarus só médias; tilápia não é fatorial completo.','modules/mod_anova_dois_fatores.R'),
('ANCOVA','truta; crabs_morfometria','Cobertura condicional','Validar bases atuais','Interação, sobreposição de covariáveis e tanque.','modules/mod_ancova.R'),
('Proporções; independência; aderência','camaroes_sexo; lagostas_kelp_sexo','Coberto; explicitar receitas','Reutilizar bases reais','Expectativas de aderência exigem justificativa; preservar frequências.','modules/mod_proporcoes.R; mod_nonparametric.R'),
('Uma e duas variâncias','Biometrias reais disponíveis','Coberto; falta roteiro de aplicação','camarao_vannamei_biometria; pinguins','Valor hipotético de variância precisa de motivo.','tests/test_parametricos_complementares.R'),
('Logística binária','Sexo em biometrias é possível; sem evento canônico explícito','Lacuna de exemplo com evento claro','rpart::kyphosis','Evento presente=1; avaliar separação, não criar classe por limiar só para preencher lacuna.','modules/mod_regression.R; registro_execucoes.R'),
('Poisson; binomial negativa','doubs riqueza; recifes com dependência','Lacuna de exemplo com esforço','MASS::ships; InsectSprays','Seis células sem exposição em ships; filtro só no ramo, diagnóstico antes de escolher NB.','modules/mod_regressao_contagem.R'),
('Linear simples; correlação','morfometria_barbo e biometrias','Coberto','Manter canônico atual','Não selecionar linhas para melhorar p-valores.','modules/mod_regression.R; mod_correlacao.R'),
('Potência; von Bertalanffy','bagley; cangulo; snapper; walleye','Coberto com documentação a completar','Manter indivíduos; retirar logs redundantes','Separar espécies e origem de idades.','modules/mod_nonlinear.R'),
('Logística/exponencial/polinomial/logarítmica','Poucos exemplos canônicos por curva explicitados','Ampliar curva logística','datasets::Orange','Não garantir uma base por curva sem diagnóstico; dependência por árvore.','modules/mod_nonlinear.R; mod_model_discovery.R'),
('Séries: visualizar/decompor/ACF','cpue_tubarao; aquicultura anual','CPUE exige fonte/unidade/regularidade','datasets::co2; NOAA ERSSTv6','Não concatenar barcos; oito anos anuais não cobrem sazonalidade mensal.','modules/mod_series_temporais.R'),
('PCA; HCA; heatmap','doubs, barbo, cabritos, salmouras, nutrição, QFASA, recifes','Coberto; priorizar clareza','Bases menores com dicionário claro','IDs e clusters externos fora da matriz; unidades/nomes químicos pendentes.','modules/mod_pca.R; mod_hca.R'),
('Amostragem AAS/AEP/AS','walleye; snapper; crabs','Coberto','Manter população e estratos','Não cortar para economizar espaço; hierarquia preservada.','modules/mod_sampling.R'),
('Coroplético; pontos/bolhas','aquicultura real; estações sintéticas','Pontos: alternativa metodológica real identificada','datasets::quakes','Fiji: fundo brasileiro desativado; longitude leste contínua. Base pesqueira georreferenciada ainda pendente.','inst/app/app.R; modules/mod_mapa_pontos.R'),
('Preparo: separar/empilhar/tratar','treinos; pescada bruta; pargo largo','Coberto','Manter dois sintéticos autorizados','Não corrigir previamente os defeitos que são alvo do exercício.','modules/registro_tratamentos.R; mod_arrumar.R'),
('Múltipla; previsão; k-means; PERMANOVA; raster/densidade','Opções em desenvolvimento/preparação na interface','Futuro, não módulo pronto','Registrar interesse sem ampliar pacote agora','Sem prometer integração à ClaRa.','inst/app/app.R: placeholders; ClaRa/README.md')]
for r in coverage:gap.append(r)
widths={'Comece aqui':[32,125],'Curadoria':[34,24,42,22,95,26,70,10,13,20,38],'Variáveis':[34,38,42,25,42,85,24,10,16,18],'Categorias':[34,30,45,45,25,45,85,18],'Diagnóstico':[10,34,30,90,80,72,12,12,18,20,90],'Cobertura':[44,55,58,46,90,75]}
editcols={'Curadoria':[6,7],'Variáveis':[3,4,5],'Categorias':[4,5,6]}
for s in wb:
    s.freeze_panes='B2' if s.title!='Comece aqui' else 'A2';s.sheet_view.zoomScale=80
    s.sheet_properties.pageSetUpPr.fitToPage=True
    s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3
    s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.print_title_rows='1:1'
    for j,width in enumerate(widths[s.title],1):s.column_dimensions[s.cell(1,j).column_letter].width=width
    for cell in s[1]:
        cell.fill=PatternFill('solid',fgColor='0F3B5F');cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=11);cell.alignment=Alignment(wrap_text=True,vertical='center')
        cell.comment=Comment('Levantamento em 10/10/2026. Todas as decisões são pendentes. Campos azuis editáveis.','Curadoria EAPADados')
    s.row_dimensions[1].height=42
    for row in s.iter_rows(min_row=2):
        for cell in row:
            editable=cell.column in editcols.get(s.title,[])
            cell.font=Font(name='Arial',size=10,color='0055AA' if editable else '172B3A')
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if editable:cell.fill=PatternFill('solid',fgColor='E8F5FF')
        s.row_dimensions[row[0].row].height=110 if s.title=='Curadoria' else 76 if s.title in ['Comece aqui','Diagnóstico','Cobertura'] else 48
    if s.title!='Comece aqui':
        t=Table(displayName='Tabela'+tidy(s.title).replace('_','').capitalize(),ref=s.dimensions)
        t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);s.add_table(t)
    s.auto_filter.ref=s.dimensions
for s,col in [(ws,'F'),(vs,'D'),(cs,'E')]:
    options='Pendente,Aprovar proposta,Aprovar com ajustes,Recusar,Adiar' if s==ws else 'Herdar do conjunto,Pendente,Aprovar proposta,Aprovar com ajustes,Recusar,Adiar'
    dv=DataValidation(type='list',formula1='"'+options+'"',allow_blank=False)
    dv.errorTitle='Escolha uma decisão';dv.error='Use uma opção da lista.';dv.showErrorMessage=True;dv.errorStyle='stop';dv.showDropDown=False
    dv.promptTitle='Sua decisão';dv.prompt='Escolha e salve. Ajustes exigem texto ao lado.';dv.showInputMessage=True;s.add_data_validation(dv);dv.add(f'{col}2:{col}{s.max_row}')
    for val,color in [('Aprovar proposta','D8F3DC'),('Aprovar com ajustes','FFF2CC'),('Recusar','FADBD8'),('Adiar','E5E7EB')]:
        s.conditional_formatting.add(f'{col}2:{col}{s.max_row}',CellIsRule(operator='equal',formula=['"'+val+'"'],fill=PatternFill('solid',fgColor=color)))
for idx,r in enumerate(rows,2):
    if r['natureza'].startswith('Sintético'):
        ws.cell(idx,11).fill=PatternFill('solid',fgColor='FFF2CC');ws.cell(idx,11).font=Font(name='Arial',size=10,bold=True,color='8B4500')
    if r['origem'].startswith('Pacote R:'):
        ws.cell(idx,4).comment=Comment('Pronto NA ORIGEM, conforme pedido. Tradução, recorte e incorporação sujeitos à aprovação; não equivale a integração já pronta.','Curadoria')
wb.save(OUT/'curadoria_eapadados_revisao.xlsx')
(OUT/'propostas.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
with (OUT/'curadoria_resumo.csv').open('w',encoding='utf-8-sig',newline='') as f:
    writer=csv.writer(f,delimiter=';');writer.writerow(headers)
    for r in rows:writer.writerow([r['nome'],r['origem'],r['analise'],r['situacao'],r['avaliacao'],r['decisao'],r['autor'],r['id'],r['prioridade'],r['dimensao'],r['natureza']])
manifest={}
for directory in ['data','R','man','data-raw','inst']:
    for f in (ROOT/directory).rglob('*'):
        if f.is_file():manifest[f.relative_to(ROOT).as_posix()]=hashlib.sha256(f.read_bytes()).hexdigest()
(OUT/'manifesto_sha256.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
check=load_workbook(OUT/'curadoria_eapadados_revisao.xlsx')
assert check['Curadoria'].max_row==len(rows)+1
assert len({r['id'] for r in rows})==len(rows)
assert [r['nome'] for r in rows[:43]]==list(inv)
assert all(r['situacao']=='pronto' for r in rows[-12:])
assert all(r['decisao']=='Pendente' for r in rows)
assert len(varrows)==sum(x['colunas'] for x in inv.values())
for n in inv:
    proposed=[r[2] for r in varrows if r[0]==n and r[2]]
    assert len(proposed)==len(set(proposed)),n
assert all(len(s.data_validations.dataValidation)==1 for s in [check['Curadoria'],check['Variáveis'],check['Categorias']])
assert not any(c.data_type=='f' for s in check for row in s for c in row)
print(json.dumps(dict(conjuntos_atuais=len(inv),candidatos=len(new),total=len(rows),variaveis=len(varrows),niveis_categoricos=len(catrows),originais_com_hash=len(manifest),xlsx_bytes=(OUT/'curadoria_eapadados_revisao.xlsx').stat().st_size),ensure_ascii=False))
