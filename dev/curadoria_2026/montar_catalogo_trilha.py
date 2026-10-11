"""Catálogo de consulta; não modifica dados, decisões ou código da TRILHA."""
from pathlib import Path
from collections import Counter
from datetime import date
import hashlib
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

PASTA = Path(__file__).resolve().parent
TRILHA = PASTA.parents[2] / 'TRILHA'
DATA = '10/10/2026'
itens = []

def grupo(area, fonte, linhas, status='Disponível'):
    for nome, faz, dados, cuidado in linhas:
        itens.append(dict(area=area, funcionalidade=nome, situacao=status,
                          descricao=faz, estrutura=dados, curadoria=cuidado,
                          fonte=fonte))

grupo('Planejamento observacional', 'inst/app/modules/mod_planejamento_observacional.R', [
 ('Transversal comparativo', 'Planeja a comparação de grupos preexistentes em um período.', 'Grupos; unidades independentes; resposta; possíveis subunidades.', 'Preservar grupo e unidade amostral; distinguir indivíduo de pool.'),
 ('Longitudinal comparativo', 'Planeja o acompanhamento das mesmas unidades ao longo do tempo.', 'Identificador estável; grupo; ocasião; resposta em formato longo.', 'Identificadores e ocasiões são indispensáveis; não tratar visitas como réplicas independentes.'),
 ('Estudo de gradiente', 'Organiza observações ao longo de um gradiente contínuo.', 'Estação/unidade; variável de gradiente; resposta; possíveis subunidades.', 'Manter gradiente contínuo e identidade de estação; não criar grupos arbitrários.'),
 ('Estudos de impacto: BA, CI e BACI', 'Planeja antes/depois, controle/impacto e a combinação BACI.', 'Local; condição; período; unidade; resposta; replicação espacial e temporal.', 'Buscar desenho real completo; o planejamento não garante um motor geral de análise BACI.')])
grupo('Planejamento e coleta', 'inst/app/modules/mod_monitoramento.R', [
 ('Monitoramento e séries temporais', 'Organiza o plano de acompanhamento temporal.', 'Indicador; unidade; frequência; datas; locais de acompanhamento.', 'Manter periodicidade, lacunas e mudanças no esforço de observação.')])
grupo('Planejamento e coleta', 'inst/app/modules/mod_planejamento_variaveis.R', [
 ('Variáveis e ficha de coleta', 'Define variáveis e organiza a estrutura da coleta nos fluxos de planejamento.', 'Nomes, tipos, unidades e unidade observacional.', 'A ficha deve preservar a hierarquia amostral e separar identificação de resposta.')])
grupo('Planejamento experimental', 'inst/app/modules/mod_experimental_design.R', [
 ('DIC — inteiramente casualizado', 'Planeja e casualiza tratamentos entre unidades experimentais.', 'Tratamento; repetição; unidade experimental; resposta futura.', 'Réplicas são unidades experimentais, não medições dentro de um tanque.'),
 ('DBC — blocos casualizados', 'Planeja tratamentos controlando a variação entre blocos.', 'Bloco; tratamento; unidade experimental.', 'Preservar blocos; não confundir disponibilidade do planejamento com análise específica de DBC.'),
 ('DQL — quadrado latino', 'Planeja tratamentos com controle por linhas e colunas.', 'Linha; coluna; tratamento; unidade experimental.', 'Manter as duas restrições de casualização; análise específica não aparece como módulo próprio.'),
 ('Fatorial em DIC', 'Planeja combinações de fatores e repetições.', 'Fatores; combinações; unidades experimentais.', 'Preferir cruzamento completo; não inventar combinações ausentes.'),
 ('Parcelas subdivididas', 'Planeja fatores aplicados em parcelas e subparcelas.', 'Parcela; subparcela; fatores; resposta.', 'Exige dois níveis de unidade/erro; não é sinônimo de ANOVA com subamostras.')])
grupo('Dimensionamento amostral', 'inst/app/modules/mod_n_poder.R', [
 ('Uma média: precisão', 'Dimensiona n a partir de variabilidade, confiança e margem de erro.', 'Estimativa de desvio-padrão e precisão desejada; população finita opcional.', 'Um conjunto piloto com unidade e escala documentadas é útil.'),
 ('Uma proporção: precisão', 'Dimensiona n para estimar uma proporção.', 'Proporção esperada; confiança; margem; população finita opcional.', 'Documentar sucesso, denominador e população-alvo.'),
 ('Comparação de médias: poder', 'Dimensiona duas médias por teste t ou três ou mais por ANOVA.', 'Efeito esperado; variabilidade; alfa; poder; número de grupos.', 'Dados piloto devem permitir estimar efeito e variabilidade entre unidades.'),
 ('Comparação de proporções: poder', 'Dimensiona a comparação de duas proporções.', 'Proporções esperadas; alfa; poder.', 'Escolher exemplo com eventos e denominadores claros.'),
 ('Correlação: poder', 'Dimensiona n para detectar associação por correlação.', 'Correlação esperada; alfa; poder.', 'Usar pares de medidas de unidades independentes.')])
grupo('Seleção da amostra', 'inst/app/modules/mod_sampling.R', [
 ('Amostragem aleatória simples', 'Sorteia unidades a partir de um cadastro, com semente.', 'Cadastro da população e identificador único.', 'Preservar o cadastro completo; resposta medida não é necessária.'),
 ('Amostragem estratificada proporcional', 'Distribui a amostra proporcionalmente entre estratos e sorteia unidades.', 'Identificador; estrato de cada unidade.', 'Não remover estratos raros sem avaliar a alocação.'),
 ('Amostragem sistemática', 'Seleciona unidades de uma lista ordenada com intervalo de seleção.', 'Cadastro ordenado; identificador; tamanho da amostra.', 'A ordem do cadastro faz parte do desenho.'),
 ('Amostragem por conglomerados', 'Sorteia conglomerados em um estágio ou unidades dentro deles em dois estágios.', 'Conglomerado; identificador das unidades; cadastro.', 'Manter vínculo unidade–conglomerado; não contar indivíduos como conglomerados.')])
grupo('Importação e organização', 'inst/app/app.R', [
 ('Importar CSV ou Excel', 'Lê CSV, XLSX e XLS; permite selecionar planilha e convenções de separador/decimal.', 'Arquivo tabular com cabeçalhos e tipos reconhecíveis.', 'Boa base para mostrar decimal com vírgula, cabeçalhos e dados ausentes.'),
 ('Abrir conjuntos do EAPADados', 'Seleciona conjuntos do pacote para exploração e análise.', 'Objeto disponível no pacote instalado.', 'A disponibilidade no seletor depende da versão instalada.'),
 ('Empilhar colunas', 'Converte dados de formato largo para longo.', 'Identificadores e várias colunas de medidas/ocasiões.', 'Exemplo útil para medidas repetidas; preservar os identificadores.'),
 ('Alargar planilha', 'Converte formato longo para largo.', 'Chaves de identificação; coluna de nomes; coluna de valores.', 'Conferir unicidade das combinações de chaves.'),
 ('Separar colunas', 'Desmembra informação composta em colunas.', 'Coluna textual com separador ou padrão reconhecível.', 'Manter a correspondência com o registro original.')])
grupo('Tratamento dos dados', 'inst/app/modules/registro_tratamentos.R', [
 ('Organizar variáveis', 'Seleciona, renomeia e ajusta tipos de variáveis na preparação.', 'Tabela com tipos e significados documentados.', 'Não descartar ID, bloco, unidade, tempo, esforço ou coordenadas por parecerem auxiliares.'),
 ('Tratar valores ausentes', 'Remove linhas ou imputa média, mediana, moda ou constante, conforme a opção.', 'Variável com NA e explicação de sua origem.', 'Imputação e exclusão não são limpezas neutras; dependem de aprovação na curadoria.'),
 ('Remover duplicatas', 'Mantém a primeira linha por combinação das colunas escolhidas.', 'Identificador ou chave que defina duplicação real.', 'Linhas iguais de indivíduos distintos não são duplicatas indevidas.'),
 ('Padronizar texto', 'Ajusta espaços e caixa: minúsculas, maiúsculas ou título.', 'Colunas textuais/categorias.', 'Conferir se rótulos distintos não serão fundidos inadvertidamente.'),
 ('Filtrar linhas', 'Recorta observações por condições numéricas ou categóricas.', 'Variável de filtro e critérios explícitos.', 'Registrar critério, n antes/depois e impacto na estrutura amostral.'),
 ('Sortear subamostras', 'Seleciona uma amostra de linhas, inclusive dentro de grupos.', 'Base com registros e grupos opcionais.', 'Não confundir sorteio didático com amostragem independente em dados hierárquicos.'),
 ('Calcular variáveis', 'Cria variáveis derivadas por expressões na receita de preparação.', 'Colunas necessárias ao cálculo.', 'Documentar fórmula, unidade e significado da nova variável.'),
 ('Reescalar unidades', 'Altera a escala de uma variável por conversão configurada.', 'Variável numérica; unidade original e unidade de destino.', 'Confirmar unidade original antes de converter ou renomear.'),
 ('Padronizar valores numéricos', 'Oferece z-score, centralização e escala min–max de 0 a 1.', 'Variáveis numéricas.', 'Preservar as medidas originais e explicar a escala derivada.'),
 ('Criar classes', 'Agrupa uma variável em intervalos de mesma largura ou por quantis.', 'Variável numérica.', 'Classes são derivadas; preservar a variável contínua.'),
 ('Dicotomizar', 'Cria 0/1 por limiar numérico ou seleção de categorias.', 'Variável numérica ou categórica.', 'Definir claramente o evento 1 e justificar o limiar.'),
 ('Agrupar e sumarizar', 'Produz resumos por combinações de grupos; encerra a receita dessa base.', 'Grupos; respostas numéricas; unidades identificadas.', 'Não substituir dados individuais por médias quando a análise precisa de réplicas.'),
 ('Construir tabela de contingência', 'Conta combinações categóricas e calcula percentuais por linha, coluna ou total.', 'Variáveis categóricas; unidades observacionais.', 'Conservar contagens; percentuais sozinhos não bastam para os testes.')])
grupo('Bases e reprodutibilidade', 'inst/app/modules/mod_bases_derivadas.R', [
 ('Base compartilhada e bases derivadas', 'Mantém preparação comum e ramificações com receitas específicas de tratamento.', 'Uma base de origem e receitas ordenadas.', 'Derivações partem da base compartilhada; preservam contexto de preparação.')])
grupo('Exploração e gráficos', 'inst/app/modules/mod_descrevendo_dados.R', [
 ('Panorama do conjunto', 'Apresenta dimensões, tipos e ausências para conhecer a base.', 'Tabela numérica e/ou categórica.', 'Exemplo deve permitir discutir significado das linhas e dados ausentes.'),
 ('Resumo das variáveis', 'Apresenta estatísticas descritivas, distribuições e frequências.', 'Variáveis numéricas e categóricas.', 'Incluir unidades e níveis compreensíveis.'),
 ('Relações entre variáveis', 'Explora associações e comparações conforme os tipos das variáveis.', 'Pares de variáveis e grupos opcionais.', 'Associação exploratória não define causalidade.'),
 ('Pressupostos e valores extremos', 'Ajuda a avaliar normalidade, gráficos Q–Q e observações extremas.', 'Variáveis numéricas; resíduos nos módulos de modelos.', 'Não excluir valores extremos automaticamente.'),
 ('Comparar transformações', 'Compara original, log, raiz quadrada e Box–Cox quando admissíveis.', 'Variável numérica; positividade conforme a transformação.', 'Exploração não altera a base; zeros e valores negativos restringem opções.'),
 ('Histograma e densidade', 'Mostra a distribuição de uma variável numérica.', 'Variável numérica.', 'Escolher exemplos com formas de distribuição distintas.'),
 ('Gráfico de barras', 'Mostra frequências e comparações categóricas.', 'Variável categórica e agrupamento quando pertinente.', 'Preservar níveis e denominadores.'),
 ('Comparar grupos graficamente', 'Oferece boxplot, violino, pontos e resumos com média, IC e n.', 'Resposta numérica; grupo.', 'Informar se n representa indivíduos, tanques ou outra unidade.'),
 ('Dispersão e tendência', 'Mostra pares numéricos, com reta linear, suavização LOESS ou sem tendência.', 'Duas variáveis numéricas; grupo opcional.', 'Buscar gradientes, não linearidade e amplitude útil.'),
 ('Linhas em eixo ordenado', 'Conecta medidas ao longo de um eixo ordenado.', 'Eixo ordenado; resposta; série/grupo se houver.', 'Conservar ordem e espaçamento; não confundir com análise formal temporal.'),
 ('Matriz de dispersão', 'Compara visualmente vários pares de variáveis numéricas.', 'Várias variáveis numéricas nas mesmas unidades.', 'Útil para biometria e qualidade da água.'),
 ('Mapa de calor de correlações', 'Resume correlações entre variáveis em uma matriz colorida.', 'Várias variáveis numéricas.', 'Correlações exigem observações correspondentes.')])
grupo('Frequências e proporções', 'inst/app/modules/mod_proporcoes.R', [
 ('Uma proporção', 'Teste binomial exato contra uma proporção de referência, com IC.', 'Resposta binária por unidade independente.', 'Informar evento de sucesso e hipótese de referência.'),
 ('Duas proporções independentes', 'Compara proporções de dois grupos por prop.test; a receita também inclui Fisher.', 'Resposta binária e grupo com dois níveis.', 'Pares, blocos e agrupamentos não são atendidos por este procedimento.'),
 ('Qui-quadrado de aderência', 'Compara frequências observadas com proporções esperadas especificadas.', 'Resposta categórica e hipótese de distribuição esperada.', 'A distribuição esperada precisa ter justificativa substantiva.')])
grupo('Frequências e proporções', 'inst/app/app.R', [
 ('Qui-quadrado de independência', 'Avalia associação em tabela de contingência.', 'Duas categorias por observação ou tabela longa com frequência n.', 'Não entregar apenas percentuais; conferir contagens esperadas.'),
 ('McNemar', 'Compara respostas binárias pareadas.', 'Duas respostas binárias do mesmo par/unidade.', 'Preservar correspondência antes/depois; buscar pares com mudanças em ambas as direções.')])
grupo('Testes paramétricos', 'inst/app/app.R', [
 ('Teste t: uma amostra', 'Compara uma média com valor de referência.', 'Resposta numérica de unidades independentes.', 'Documentar a referência e a unidade da resposta.'),
 ('Teste t: dois grupos independentes', 'Compara duas médias; Welch ou Student com variâncias iguais.', 'Resposta numérica; grupo com dois níveis.', 'Não converter dados pareados em independentes.'),
 ('Teste t pareado', 'Compara duas medidas correspondentes.', 'Duas medidas numéricas por unidade/par.', 'ID e pareamento não podem se perder no recorte.'),
 ('ANOVA de um fator', 'Compara grupos; opção clássica/Tukey ou Welch/Games–Howell.', 'Resposta numérica; fator; réplicas independentes.', 'Médias, DP e n não substituem as observações para este fluxo.'),
 ('ANOVA com subamostras', 'Separa unidades e subunidades, com efeito aleatório da unidade; contrasta abordagens didáticas.', 'Tratamento; unidade experimental; subunidade; resposta.', 'Prioridade de curadoria: dados hierárquicos reais e documentação da unidade experimental.'),
 ('ANOVA de medidas repetidas', 'Compara momentos das mesmas unidades no fluxo de um fator temporal.', 'ID; momento; resposta em formato longo; grade completa.', 'Buscar seguimento real; não assumir suporte geral a desenhos longitudinais mistos.'),
 ('ANOVA de dois fatores', 'Estima efeitos de dois fatores e sua interação.', 'Dois fatores; resposta; réplicas nas combinações.', 'Preferir fatorial completo; verificar células vazias e controles externos ao fatorial.'),
 ('ANCOVA', 'Compara grupos ajustando por covariável numérica.', 'Fator; covariável contínua; resposta numérica.', 'Buscar sobreposição da covariável entre grupos e relação interpretável.'),
 ('Uma variância: qui-quadrado', 'Testa uma variância contra valor de referência.', 'Resposta numérica e variância de referência.', 'Referência deve usar a unidade ao quadrado; pressuposto de normalidade.'),
 ('Duas variâncias: teste F', 'Compara variâncias de dois grupos.', 'Resposta numérica; dois grupos independentes.', 'Exemplo deve permitir discutir sensibilidade à não normalidade.')])
grupo('Testes não paramétricos', 'inst/app/modules/mod_nonparametric.R', [
 ('Mann–Whitney', 'Compara dois grupos independentes por postos.', 'Resposta numérica/ordenável; dois grupos.', 'Preservar observações individuais e empates.'),
 ('Wilcoxon pareado', 'Testa diferenças entre pares por postos sinalizados.', 'Duas medidas correspondentes por unidade.', 'Manter pares e diferenças nulas; não reordenar as colunas separadamente.'),
 ('Kruskal–Wallis e Dunn', 'Compara vários grupos independentes e oferece pós-teste de Dunn com ajustes.', 'Resposta ordenável; fator com três ou mais grupos.', 'Pós-teste depende dos pacotes auxiliares disponíveis.')])
grupo('Testes não paramétricos', 'inst/app/app.R', [
 ('Friedman', 'Compara condições pareadas; comparações Wilcoxon com ajuste de Holm.', 'Unidade/bloco; condição; resposta; grade completa.', 'Prioridade: medidas repetidas reais sem perda de identificação.')])
grupo('Correlação e regressão', 'inst/app/modules/mod_correlacao.R', [
 ('Correlação de Pearson, Spearman e Kendall', 'Quantifica associação entre duas variáveis com alternativas paramétrica e por postos.', 'Duas variáveis correspondentes por unidade.', 'Incluir exemplos com empates, relação monotônica e valores extremos.')])
grupo('Correlação e regressão', 'inst/app/modules/mod_regression.R', [
 ('Regressão linear simples', 'Ajusta resposta numérica por um preditor numérico e apresenta diagnósticos.', 'Resposta contínua; preditor numérico.', 'Priorizar amplitude do preditor e unidades independentes.')])
grupo('Correlação e regressão', 'inst/app/app.R', [
 ('Regressão logística binária', 'Modela probabilidade de resposta 0/1.', 'Resposta binária; preditor(es) admitidos pelo módulo.', 'Conferir sucesso/fracasso e evitar separação completa.')])
grupo('Modelos de contagem', 'inst/app/modules/mod_regressao_contagem.R', [
 ('Regressão de Poisson', 'Modela contagens; permite múltiplos preditores e exposição positiva como offset.', 'Contagem inteira não negativa; preditores; esforço/exposição opcional.', 'Preservar zeros e esforço; CPUE contínua não substitui a contagem com offset.'),
 ('Regressão binomial negativa', 'Modela contagens com sobredispersão; aceita exposição como offset.', 'Contagem; preditores; exposição opcional.', 'Buscar sobredispersão real; não produzir contagens artificiais por arredondamento.')])
grupo('Curvas e crescimento', 'inst/app/app.R', [
 ('Descobrindo o modelo', 'Apoia a exploração de formas de relação e escolha de curva.', 'Pares de variáveis numéricas.', 'Escolher exemplo com mecanismo e domínio conhecidos.')])
grupo('Curvas e crescimento', 'inst/app/modules/mod_nonlinear.R', [
 ('Modelo exponencial', 'Ajusta Y = a × exp(bX).', 'X e Y numéricos; domínio apropriado.', 'Dados reais com amplitude suficiente para identificar a curvatura.'),
 ('Modelo de potência / alometria', 'Ajusta Y = a × X^b; inclui relação peso–comprimento.', 'Comprimento e peso, ou outro par positivo.', 'Confirmar unidades, espécies, estágios e possíveis grupos.'),
 ('Crescimento de von Bertalanffy', 'Ajusta curva de crescimento por idade/tempo.', 'Idade/tempo e tamanho; cobertura de idades.', 'Buscar medidas reais; poucos indivíduos jovens não identificam bem a assíntota.'),
 ('Modelo logístico de crescimento', 'Ajusta curva sigmoidal contínua.', 'Tempo e tamanho/biomassa com fases da curva.', 'É crescimento logístico, diferente da regressão logística de resposta binária.'),
 ('Modelo polinomial quadrático', 'Ajusta relação com termo de segundo grau.', 'X e Y numéricos.', 'Não extrapolar além do domínio observado.'),
 ('Modelo logarítmico', 'Ajusta Y = a + b × ln(X).', 'X positivo; Y numérico.', 'Não deslocar zeros arbitrariamente para permitir o ajuste.')])
grupo('Séries temporais', 'inst/app/modules/mod_series_temporais.R', [
 ('Visualizar e suavizar séries', 'Mostra a evolução, agrega por médias mensais/anuais e aplica média móvel.', 'Data; resposta numérica; série identificada.', 'Conservar datas, lacunas, periodicidade e significado da agregação.'),
 ('Decomposição aditiva', 'Separa tendência, sazonalidade e resíduos de séries mensais.', 'Série mensal com pelo menos 24 meses.', 'O módulo usa interpolação temporária para lacunas; não altera a base original.'),
 ('Autocorrelação — ACF', 'Examina dependência entre observações defasadas.', 'Série temporal ordenada e espaçamento regular.', 'Evitar tratar as observações como unidades independentes; conferir tratamento das lacunas.')])
grupo('Análise multivariada', 'inst/app/modules/mod_pca.R', [
 ('PCA — componentes principais', 'Resume várias variáveis numéricas; padronização, componentes e biplots; opção de imputação.', 'Matriz de unidades × variáveis numéricas; grupo opcional.', 'Conferir escalas, ausências e unidade da linha; imputação via missMDA é opcional.')])
grupo('Análise multivariada', 'inst/app/modules/mod_hca.R', [
 ('Agrupamento hierárquico — HCA', 'Agrupa linhas; distâncias euclidiana, Manhattan e binária/Jaccard; ligações Ward, completa, simples e média.', 'Matriz numérica; presença/ausência para distância binária.', 'A orientação da matriz define quem é agrupado; Jaccard/binária não usa z-score.')])
grupo('Mapas', 'inst/app/modules/mod_mapa_pontos.R', [
 ('Mapa de pontos / onde amostrar', 'Localiza estações e unidades por latitude/longitude.', 'Latitude; longitude; identificação do local.', 'Conferir datum, sinal e convenção de longitude; sorteio espacial ainda não é recurso disponível.'),
 ('Mapa de bolhas proporcionais', 'Representa magnitude por tamanho de bolha; permite grupos e rótulos.', 'Coordenadas e magnitude numérica não negativa.', 'Prioridade: dados reais georreferenciados com metadados e variável mapeável.')])
grupo('Mapas', 'inst/app/modules/mod_mapa.R', [
 ('Mapa coroplético por UF', 'Colore unidades federativas do Brasil por indicador.', 'UF reconhecível e valor por estado.', 'Compatibilizar códigos de UF e ano; evitar misturar total e taxa.')])
grupo('Laboratório de conceitos', 'inst/app/modules/mod_laboratorio.R', [
 ('Teorema do Limite Central: médias', 'Simula distribuições de médias amostrais.', 'Parâmetros da simulação; não exige conjunto empírico.', 'Simulação didática interna não corresponde à incorporação de novo conjunto sintético.'),
 ('Simuladores de teste t e ANOVA', 'Exploram resultados por simulação nos painéis de testes.', 'Parâmetros de grupos, variabilidade e tamanho amostral.', 'Não exige novos conjuntos de dados para o pacote.')])
grupo('Comunicação dos resultados', 'inst/app/modules/mod_comunicacao.R', [
 ('Projeto R e relatório reproduzível', 'Seleciona execuções integradas, organiza texto/componentes e exporta projeto ZIP com código e Quarto.', 'Análises registradas e suas bases/receitas.', 'Word e HTML são gerados ao renderizar no RStudio; nem todo módulo está integrado ao registro.')])
grupo('Recursos em preparação', 'inst/app/app.R', [
 ('Regressão linear múltipla', 'Entrada de menu ainda apresenta módulo em desenvolvimento.', 'Resposta numérica e múltiplos preditores, quando implementado.', 'Distinguir da regressão de contagem, que já aceita múltiplos preditores.'),
 ('Previsão de séries temporais', 'Tela introdutória em preparação; não há fluxo de previsão pronto.', 'Série regular suficientemente longa, quando implementado.', 'Não escolher dados pressupondo ARIMA/ETS já disponíveis.'),
 ('K-means', 'Entrada de menu em preparação.', 'Matriz numérica, quando implementado.', 'Pode ser atendido futuramente por dados já adequados à PCA/HCA.'),
 ('PERMANOVA', 'Entrada de menu em preparação.', 'Matriz multivariada e desenho/grupos, quando implementado.', 'Conservar desenho amostral e matriz original para uso futuro.'),
 ('Densidade de ocorrências', 'Tela de mapa em preparação; arquivo de módulo existente não equivale a integração ativa.', 'Ocorrências georreferenciadas, quando implementado.', 'Não classificar o recurso como operacional apenas por existir arquivo de código.'),
 ('Dependência espacial', 'Tela em preparação.', 'Dados espaciais e definição de vizinhança, quando implementado.', 'Não há fluxo ativo de análise espacial nesse item.'),
 ('Ambiente nos pontos', 'Tela em preparação.', 'Pontos e dados ambientais/raster, quando implementado.', 'Extração ambiental/raster não está disponível nesse fluxo de menu.')], 'Em preparação')
grupo('Laboratório em preparação', 'inst/app/modules/mod_laboratorio.R', [
 ('LLN, IC e TLC para proporções', 'Lei dos Grandes Números, cobertura de IC e TLC de proporções são telas em construção.', 'Simulações futuras.', 'Não exigem novos conjuntos empíricos.'),
 ('Distribuições de probabilidade', 'Visualizadores Normal, binomial, Poisson, t, F e qui-quadrado estão em construção neste laboratório.', 'Parâmetros de distribuição.', 'Isto não impede o uso dessas distribuições nos testes já implementados.'),
 ('P-valor, permutação, erros e poder', 'Visualizadores conceituais em construção.', 'Simulações futuras.', 'Dimensionamento amostral por poder já existe em outro módulo.'),
 ('Simuladores qui-quadrado e correlação', 'Abas em construção; simuladores t e ANOVA já estão disponíveis.', 'Simulações futuras.', 'Testes reais de qui-quadrado e correlação já existem em outros menus.'),
 ('Anscombe/Datasaurus, teste de fila e dispersão', 'Painéis conceituais de visualização em construção.', 'Exemplos pedagógicos futuros.', 'Não confundir presença no catálogo com ferramenta executável.')], 'Em preparação')

prioridades = [
 ('Pares e seguimento', 't pareado, Wilcoxon, McNemar, Friedman e medidas repetidas', 'ID estável, ocasiões/condições e respostas correspondentes; binárias e contínuas.'),
 ('Hierarquia amostral', 'ANOVA com subamostras', 'Tratamento, unidade experimental e subunidades reais; réplicas independentes identificadas.'),
 ('Dois fatores e covariáveis', 'ANOVA de dois fatores e ANCOVA', 'Fatorial com replicação; ou grupos com covariável contínua e sobreposição de valores.'),
 ('Eventos e esforço', 'Logística binária, Poisson e binomial negativa', 'Eventos reais 0/1; contagens inteiras com zeros, exposição e preditores; sobredispersão documentada.'),
 ('Tempo e sazonalidade', 'Séries, decomposição e ACF', 'Datas regulares, pelo menos dois ciclos anuais mensais, lacunas e esforço documentados.'),
 ('Múltiplas variáveis', 'PCA e HCA', 'Unidades × medidas numéricas; ou presença/ausência com orientação explícita.'),
 ('Espaço', 'Pontos, bolhas e coroplético', 'Coordenadas verificadas com indicador; ou UF e indicador comparável por estado.'),
 ('Cadastros e delineamentos', 'Amostragem e planejamento', 'Lista completa de unidades, estratos, conglomerados, blocos e níveis de casualização.'),
]

def main():
    for i, item in enumerate(itens, 1):
        item['id'] = f'TR{i:03d}'
        assert (TRILHA / item['fonte']).is_file(), item['fonte']
    saidas = [PASTA/'funcionalidades_trilha.md', PASTA/'funcionalidades_trilha.xlsx', PASTA/'funcionalidades_trilha.json']
    if any(p.exists() for p in saidas):
        raise SystemExit('Arquivos já existem; preservar e revisar antes de atualizar.')
    notas = [
        f'Consulta ao menu e ao código local da TRILHA em {DATA}.',
        'Disponível significa conectado ao menu e com implementação identificada. Este levantamento não é uma homologação de todos os fluxos pela interface nem teste de todas as bases.',
        'Em preparação significa tela informativa/placeholder no fluxo atual. Código legado ou arquivo isolado não foi contado como funcionalidade disponível.',
        'Planejamento de DBC, DQL, parcelas subdivididas ou BACI não implica disponibilidade de um motor específico de análise para todo esse desenho.',
        'Diagnósticos, pós-testes, gráficos e exportações variam por módulo e por dependências instaladas. Não há garantia de que todos os módulos participem do Projeto R integrado.',
        'A planilha é uma referência de funcionalidades. As decisões sobre os conjuntos continuam exclusivamente em curadoria_eapadados_revisao.xlsx; nenhum dado foi alterado.',
        'A estrutura indicada orienta a busca de exemplos; não constitui aprovação para remover linhas, criar dados sintéticos, imputar, recodificar ou incorporar conjuntos.',
    ]
    linhas = ['# Funcionalidades da TRILHA para orientar a curadoria', '', *[f'- {n}' for n in notas], '',
              '## O que priorizar na escolha de conjuntos', '',
              '| Necessidade | Recursos atendidos | Estrutura que deve estar presente |', '|---|---|---|']
    linhas += ['| '+' | '.join(p)+' |' for p in prioridades]
    area_anterior = None
    for item in itens:
        if item['area'] != area_anterior:
            area_anterior = item['area']
            linhas += ['', '## '+area_anterior, '', '| Recurso | Situação | O que faz | Dados necessários | Atenção na curadoria |', '|---|---|---|---|---|']
        linhas.append('| '+' | '.join(item[k].replace('|', '/') for k in ['funcionalidade','situacao','descricao','estrutura','curadoria'])+' |')
    linhas += ['', '## Fontes locais', '',
               'A situação de cada recurso foi confrontada com o menu atual em [app.R]('+str(TRILHA/'inst/app/app.R').replace('\\','/')+'). Os registros de tratamento e de execução ajudam a distinguir o que é calculado e o que entra no relatório integrado.', '']
    for fonte in sorted(set(i['fonte'] for i in itens) | {'inst/app/modules/registro_execucoes.R', 'R/descrevendo_dados.R'}):
        linhas += [f'- [{fonte}]({(TRILHA/fonte).as_posix()})']
    linhas += ['', '## Observações de interpretação', '',
        '- O teste de uma proporção usa `binom.test`; duas proporções usam `prop.test`, e a receita também contém `fisher.test`. Fisher não foi contado como uma entrada independente de menu.',
        '- A comparação exploratória de transformações inclui Box–Cox, respeitando o domínio dos dados. Ela não aplica a transformação à base.',
        '- Polinomial é a opção quadrática identificada no módulo de curvas; não foi generalizada para seleção livre de qualquer grau.',
        '- A ANOVA de medidas repetidas oferece o fluxo simples por momento e unidade; não foi tratada como solução geral para modelos longitudinais, dados faltantes ou estruturas de covariância.',
        '- Mapas de pontos não incluem sorteio espacial de estações. Densidade, dependência espacial, extração ambiental e raster não foram listados como fluxos operacionais.',
        '- As simulações didáticas internas do laboratório não propõem a entrada de outros conjuntos sintéticos no EAPADados.',
        '- Antes de uma incorporação aprovada, testar o conjunto no módulo pretendido. A adequação estatística depende do desenho, e não apenas dos tipos das colunas.', '']
    saidas[0].write_text('\n'.join(linhas), encoding='utf-8')
    saidas[2].write_text(json.dumps({'data':DATA,'criterio':notas,'itens':itens,'prioridades':prioridades}, ensure_ascii=False, indent=2), encoding='utf-8')
    wb = Workbook()
    inicio = wb.active
    inicio.title = 'Como ler'
    inicio.append(['FUNCIONALIDADES DA TRILHA', 'Referência para escolher conjuntos'])
    for n in notas:
        inicio.append(['Orientação', n])
    counts = Counter(i['situacao'] for i in itens)
    inicio.append(['Itens documentados', len(itens)])
    for s, n in counts.items():
        inicio.append([s, n])
    inicio.append(['Planilha de decisões', 'curadoria_eapadados_revisao.xlsx'])
    inicio.cell(inicio.max_row, 2).hyperlink = 'curadoria_eapadados_revisao.xlsx'
    ws = wb.create_sheet('Funcionalidades')
    cabecalhos = ['ID', 'Área', 'Funcionalidade', 'Situação', 'O que faz', 'Dados necessários', 'Atenção na curadoria', 'Fonte no código']
    ws.append(cabecalhos)
    for i in itens:
        ws.append([i[k] for k in ['id','area','funcionalidade','situacao','descricao','estrutura','curadoria','fonte']])
        ws.cell(ws.max_row, 8).hyperlink = (TRILHA/i['fonte']).as_uri()
    tab = Table(displayName='CatalogoTrilha', ref=f'A1:H{ws.max_row}')
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tab)
    ps = wb.create_sheet('Escolha dos dados')
    ps.append(['Necessidade', 'Recursos atendidos', 'Estrutura que deve estar presente'])
    for p in prioridades:
        ps.append(p)
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            for cell in row:
                cell.font = Font(name='Arial', size=11, color='203746')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='123D52')
            cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        sheet.row_dimensions[1].height = 32
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:1'
    for col, largura in zip('ABCDEFGH', [10,28,39,20,67,60,75,60]):
        ws.column_dimensions[col].width = largura
    for row in range(2, ws.max_row+1):
        ws.row_dimensions[row].height = 68
        cel = ws.cell(row,4)
        cel.fill = PatternFill('solid', fgColor='D9EFE7' if cel.value=='Disponível' else 'FFF0CC')
        cel.font = Font(name='Arial',size=11,bold=True,color='245B48' if cel.value=='Disponível' else '7B5100')
    inicio.column_dimensions['A'].width = 32
    inicio.column_dimensions['B'].width = 115
    for row in range(2,inicio.max_row+1):
        inicio.row_dimensions[row].height = 48
    for col, largura in zip('ABC',[33,65,100]):
        ps.column_dimensions[col].width = largura
    for row in range(2,ps.max_row+1):
        ps.row_dimensions[row].height = 57
    ps.auto_filter.ref = ps.dimensions
    wb.save(saidas[1])
    conferido = load_workbook(saidas[1])
    assert conferido['Funcionalidades'].max_row == len(itens)+1
    assert all(c.data_type != 'f' for s in conferido for row in s for c in row)
    assert len(set(i['id'] for i in itens)) == len(itens)
    print(json.dumps({'arquivos':[str(p) for p in saidas], 'itens':len(itens), 'situacao':dict(counts)},ensure_ascii=False))

if __name__ == '__main__':
    main()
