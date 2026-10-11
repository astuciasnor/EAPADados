"""Lê a revisão salva e gera um relatório JSON. Nunca modifica conjuntos de dados."""
from pathlib import Path
from collections import Counter
import json, hashlib, sys
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
source=Path(sys.argv[1]).resolve() if len(sys.argv)>1 else OUT/'curadoria_eapadados_revisao.xlsx'
wb=load_workbook(source,data_only=False)
proposals=json.loads((OUT/'propostas.json').read_text(encoding='utf-8'))
byid={r['id']:r for r in proposals}
allowed={'Pendente','Aprovar proposta','Aprovar com ajustes','Recusar','Adiar'}
errors=[];main=[];seen=set()
for row in wb['Curadoria'].iter_rows(min_row=2,values_only=True):
    name,origin,analysis,status,evaluation,decision,author,rid,*_=row
    if rid not in byid or rid in seen:
        errors.append(f'ID ausente, desconhecido ou duplicado: {rid}');continue
    seen.add(rid);original=byid[rid]
    if name!=original['nome']:errors.append(f'{rid}: nome identificador alterado; manter nome original.')
    if decision not in allowed:errors.append(f'{rid}: decisão inválida: {decision}')
    if decision=='Aprovar com ajustes' and not str(author or '').strip():errors.append(f'{rid}: falta texto do ajuste.')
    if evaluation!=original['avaliacao']:
        errors.append(f'{rid}: proposta original alterada. Coloque o ajuste em Avaliação final do autor.')
    main.append(dict(id=rid,nome=name,decisao=decision,acao_proposta=original['acao'],avaliacao_original=evaluation,avaliacao_autor=author or '',restricao=original['restricao']))
if seen!=set(byid):errors.append('A tabela principal não contém todos os IDs originais.')
byname={r['nome']:r for r in main}
details=[]
for sheet,dcol,acol,pcol,pendcol in [('Variáveis',3,4,2,9),('Categorias',4,5,3,7)]:
    for lineno,values in enumerate(wb[sheet].iter_rows(min_row=2,values_only=True),2):
        name=values[0];decision=values[dcol];author=values[acol];proposed=values[pcol]
        if name not in byname:
            errors.append(f'{sheet}!{lineno}: conjunto desconhecido.');continue
        if decision not in allowed|{'Herdar do conjunto'}:errors.append(f'{sheet}!{lineno}: decisão inválida.')
        if decision=='Aprovar com ajustes' and not str(author or '').strip():errors.append(f'{sheet}!{lineno}: ajuste vazio.')
        inherited=decision=='Herdar do conjunto'
        effective=byname[name]['decisao'] if inherited else decision
        pending=values[pendcol]=='Sim' or not str(proposed or '').strip()
        details.append(dict(aba=sheet,linha=lineno,conjunto=name,identificador_atual=list(values[1:pcol]),proposto=proposed,decisao=decision,decisao_efetiva=effective,ajuste_autor=author or '',pendencia_tecnica=pending,revisao_textual_necessaria=effective=='Aprovar com ajustes',liberacao_da_base=byname[name]['decisao'] in {'Aprovar proposta','Aprovar com ajustes'}))
manifest=json.loads((OUT/'manifesto_sha256.json').read_text(encoding='utf-8'))
changed=[rel for rel,h in manifest.items() if not (ROOT/rel).is_file() or hashlib.sha256((ROOT/rel).read_bytes()).hexdigest()!=h]
if changed:errors.append('Arquivos de origem mudaram desde o inventário: refazer comparação antes de execução.')
result=dict(tipo='relatorio_de_decisoes_sem_execucao',planilha=str(source),contagem=dict(Counter(r['decisao'] for r in main)),erros=errors,originais_alterados=changed,conjuntos=main,de_para=details,aviso='Este relatório não aplica alterações. Aprovação com ajustes, pendência técnica, conflito entre abas ou licença indefinida exigem leitura e resolução antes de execução. Decisão individual não libera base recusada/adiada/pendente.')
result['documentacao_obrigatoria_por_conjunto_modificado']={
    'exigencia':'Toda adaptação efetivamente aplicada exige atualização da documentação antes de ser considerada concluída.',
    'itens':['Declarar que os dados foram adaptados da fonte original.',
             'Identificar autores/título/ano/DOI ou URL; para pacote R, pacote/objeto/versão; preservar créditos e licença disponível.',
             'Listar colunas removidas e motivos.',
             'Descrever filtros de linhas, critérios e n antes/depois; declarar preservação quando não houver exclusão.',
             'Registrar de-para dos nomes de variáveis e níveis dos fatores.',
             'Descrever mudanças de formato, unidades e cálculos derivados, quando houver.',
             'Apontar script reprodutível; atualizar fonte roxygen, man, dimensões, dicionário e exemplos.',
             'Descrever somente alterações realizadas, preservando histórico relevante; propostas pendentes não são adaptações concluídas.']}
(OUT/'decisoes_lidas.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(contagem=result['contagem'],erros=errors,originais_alterados=len(changed),saida=str(OUT/'decisoes_lidas.json')),ensure_ascii=False))
raise SystemExit(1 if errors else 0)
