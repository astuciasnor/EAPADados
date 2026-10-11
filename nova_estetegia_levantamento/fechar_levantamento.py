"""Consolida propostas; não adapta, remove ou incorpora dados."""
import json, hashlib
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
PKG = ROOT.parent
TECH = ROOT / 'levantamento_tecnico'
OLD = PKG / 'dev' / 'curadoria_2026'
def read(path): return json.loads(path.read_text(encoding='utf-8-sig'))
def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
protect = [ROOT / n for n in ['curadoria_eapadados_revisao1.xlsx', 'levantamento_conjuntos_trilha.xlsx', 'tilapia_sensorial_organizada.xlsx', 'enobieta_baci_organizado.xlsx', 'Puromycin_documentado.xlsx']]
hashes = {str(p):sha(p) for p in protect}
manifest = read(OLD / 'manifesto_sha256.json')
changes = [n for n,h in manifest.items() if not (PKG/n).is_file() or sha(PKG/n)!=h]
if changes: raise RuntimeError(f'Arquivos originais divergentes do manifesto: {changes}')
proposals = {x['nome']:x for x in read(OLD/'propostas.json')}
invent = read(OLD/'inventario_tecnico.json')
measured = read(TECH/'medicao_conjuntos.json')['conjuntos']
extra = read(TECH/'medicao_fontes_extras.json')
review = load_workbook(ROOT/'curadoria_eapadados_revisao1.xlsx',read_only=True,data_only=True)
decisions = {r[0]: {'decisao':r[5], 'comentario':r[6] or ''} for r in list(review['Curadoria'].values)[1:]}
plan = load_workbook(ROOT/'levantamento_conjuntos_trilha.xlsx',read_only=True,data_only=True)
uses = {}
for r in list(plan['levantamento_conjuntos'].values)[1:]:
    if r[3] and r[4]: uses.setdefault(f'{r[3]}::{r[4]}',[]).append(r[2])

synthetic = {'artemia','captura_petrechos','estacoes_ictiofauna','ocorrencias_peixes','sst_costa_norte','tilapia_crescimento'}
support = {'dicionario_variaveis_amostragem','referencias_dados_amostragem','qfasa_acidos_graxos'}
chosen_r = {
    'FSAdata::WhiteGrunt2','FSAdata::Cabezon','FSAdata::BGHRsample','FSAdata::BGHRfish',
    'datasets::PlantGrowth','datasets::npk','MASS::oats','agridat::cochran.latin',
    'datasets::CO2','MASS::bacteria','datasets::InsectSprays','MASS::ships',
    'MASS::Insurance','MASS::quine','datasets::quakes','astsa::rec',
    'datasets::Puromycin','datasets::cars','MASS::wtloss','datasets::Orange','FSAdata::RuffeSLRH92',
}
rows=[]
def add(name,origin,analysis,status,assessment,role,size,dimension='',nature='',source='',old_id=''):
    d=decisions.get(name,{'decisao':'Pendente','comentario':''})
    rows.append(dict(nome=name,origem=origin,analise=analysis,situacao=status,avaliacao=assessment,
        papel=role,bytes=size,dimensao=dimension,natureza=nature,fonte=source,
        decisao_anterior=d['decisao'],comentario_anterior=d['comentario'],decisao_fase2='Pendente',
        avaliacao_autor='',id_anterior=old_id))

for name,v in invent.items():
    p=proposals[name]
    role='Sintético extra - avaliar retirada' if name in synthetic else 'Apoio - não contar como base analítica' if name in support else 'Núcleo proposto'
    note=p['avaliacao']
    if name=='artemia': note+=' Substituto real localizado: Artemia_growth_sterols, experimento em béqueres de 4 L; alimentação com alga enriquecida com fitosteróis. Avaliar se atende ao pedido do autor.'
    add(name,'EAPADados',p['analise'],p['situacao'],note,role,(PKG/'data'/v['arquivo']).stat().st_size,
        f"{v['linhas']} × {v['colunas']}",p.get('natureza',''),p['fonte_local'],p['id'])

def bytes_extra(names): return sum(extra[n]['bytes_xz'] for n in names)
add('whitefish_salinidade','Artigo / Dryad','Subamostras por tanque; crescimento; salinidade','precisa de ajuste',
    'Quatro CSVs obtidos. Preservar tanque de origem/atual e dia; indivíduos não marcados, sem pareamento individual. Conferir alocação e replicação antes de definir a análise.',
    'Núcleo proposto',bytes_extra([f'File_{a}.csv' for a in 'ABCD']),'4 tabelas','Real - estrutura a validar','https://doi.org/10.5061/dryad.bg79cnpph','N001')
meso=[n for n in extra if n.endswith('_Dryad.xlsx')]
add('mesocosmos_secagem','Artigo / Dryad','Fatorial; respostas ecológicas; medidas ao longo do tempo','precisa de ajuste',
    'Seis tabelas obtidas; uma entrada de estudo, não seis experimentos independentes. Preservar tanque, fatores e ocasiões; documentar exclusões já feitas na fonte.',
    'Núcleo proposto',bytes_extra(meso),'6 tabelas','Real - estrutura a validar','https://doi.org/10.5061/dryad.hx3ffbgkn','N002')
add('artemia_fitoesterois','Base aberta / BCO-DMO','Alimentação; crescimento e biomassa; comparação de tratamentos','precisa de ajuste',
    '69 registros; quatro experimentos, três béqueres de 4 L por tratamento. Confirmar níveis, unidades e variável principal. Não juntar experimentos como se fossem réplicas intercambiáveis.',
    'Núcleo proposto',extra['Artemia_growth_sterols.csv']['bytes_xz'],'69 × 5','Real','https://www.bco-dmo.org/dataset/724179/description')
newfiles=[
    ('tilapia_sensorial','tilapia_sensorial_organizada.xlsx','Friedman; notas pareadas; preparo largo/longo','122 × 7',
     '122 avaliadores, três tratamentos, aroma e sabor; analisar atributos separadamente. Confirmar âncoras e sentido da escala -3 a +3. Não há notas ausentes.', 'https://doi.org/10.17632/knsz9r4vby.5'),
    ('enobieta_baci','enobieta_baci_organizado.xlsx','Planejamento BACI; subamostras; comunidades; comparação temporal','135 macroinvertebrados + 2.700 substrato',
     'Uma intervenção, quatro controles e quatro estações a jusante; uma campanha por fase. Surber/partículas são subamostras. Reservatório tem apenas depois. Útil também para atividade de planejamento prospectivo; não equivale a intervenções independentes replicadas.', 'https://doi.org/10.5061/dryad.w3r2280ws'),
]
for name,file,analysis,dim,note,url in newfiles:
    add(name,'Artigo / dados fornecidos pelo autor',analysis,'precisa de ajuste',note,'Núcleo proposto',(ROOT/file).stat().st_size,dim,'Real - conferir fonte e protocolo',url)

special_uses={
    'agridat::cochran.latin':'Quadrado latino', 'FSAdata::BGHRfish':'Conglomerados; junção com amostras',
    'FSAdata::BGHRsample':'Conglomerados; esforço; junção com peixes', 'MASS::wtloss':'Curva exponencial',
    'datasets::Orange':'Crescimento logístico; repetição por árvore', 'astsa::rec':'Série mensal; decomposição; ACF',
    'MASS::bacteria':'McNemar; longitudinal binário', 'MASS::ships':'Poisson; exposição/offset',
    'MASS::Insurance':'Contagem com exposição/offset', 'MASS::quine':'Binomial negativa; superdispersão',
    'datasets::cars':'Modelo quadrático; comparação de curvas', 'datasets::CO2':'Medidas repetidas; resposta por planta',
    'datasets::InsectSprays':'Contagem; Poisson/NB; Kruskal-Wallis; variâncias',
}
for name,v in measured.items():
    p=proposals.get(name,{})
    analysis=p.get('analise') or special_uses.get(name) or '; '.join(uses.get(name,[])) or 'Alternativa de estrutura; consultar documentação da fonte'
    note=p.get('avaliacao') or 'Objeto obtido e medido; conferir protocolo, fonte/licença, nomes e encaixe antes da incorporação.'
    if v.get('alerta_origem'): note+=' Alerta automático de procedência (não é conclusão): '+v['alerta_origem']
    if name=='MASS::bacteria': note+=' Validação anterior na TRILHA: 44 pares completos nas semanas 0/11; perdas documentadas; McNemar executado. Não imputar pares para completar a tabela.'
    if name=='ClustBlock::fish': note+=' Dados CATA não equivalem automaticamente a notas ordinais pareadas para Friedman.'
    if name=='palmerpenguins::penguins': note+=' Sobrepõe a fonte de pinguins já no EAPADados; não contar como novo estudo independente.'
    if name=='ade4::doubs': note+=' Inclui a fonte do doubs_ambiente e tabelas de peixes/coordenadas locais; sobreposição parcial.'
    add(name,'Pacote R: '+v['pacote'],analysis,'pronto',note,'Núcleo proposto' if name in chosen_r else 'Reserva',v['bytes_rda_xz'],
        f"{v.get('linhas','?')} × {v.get('colunas','?')}",p.get('natureza','Fonte R - procedência/licença a revisar'),v['arquivo'],p.get('id',''))
for name,analysis,note in [
    ('FSAdata::YPerchRL','Monitoramento anual; recrutamento','Série anual de 1942–1960 obtida. Não usar como série sazonal mensal.'),
    ('FESta::StdEffort_dataset','Captura; esforço; padronização de esforço','Fonte empírica identificada. Captura em toneladas, não uma contagem inteira para Poisson.'),
]:
    v=extra[name]
    add(name,'Pacote R: '+name.split('::')[0],analysis,'pronto',note,'Reserva',v['bytes_xz'],f"{v['linhas']} × {v['colunas']}",'Real - conferir adequação',str(TECH/'fontes_extraidas'/name.split('::')[0]))
add('datasets::Puromycin','Pacote R: datasets / planilha fornecida','Logarítmico como aproximação; saturação Michaelis-Menten','pronto',
    '23 registros, seis concentrações por condição. R² log = 0,9652/0,9613, conferidos. Saturação é o uso principal na ajuda do R; log não tem assíntota. Um experimento por condição; não presumir réplicas biológicas. Preservar planilha; propor excluir apenas rownames da futura versão analítica.',
    'Núcleo proposto',(ROOT/'Puromycin_documentado.xlsx').stat().st_size,'23 × 4','Real - limite de replicação','https://stat.ethz.ch/R-manual/R-devel/library/datasets/html/Puromycin.html')
for i,r in enumerate(rows,1): r['id']=f'L{i:03}'
assert len(rows)==105 and sum(r['papel']=='Núcleo proposto' for r in rows)==60
assert sum(r['papel']=='Reserva' for r in rows)==36
total=sum(r['bytes'] for r in rows)
assert total==841447, total
core=[r for r in rows if r['papel']=='Núcleo proposto']
corebytes=sum(r['bytes'] for r in core)

coverage_groups=[
 ('TR001 TR002 TR003 TR004 TR005 TR006','Planejamento observacional e coleta','pinguins; truta_riacho_crescimento; doubs_ambiente; enobieta_baci; astsa::rec; FSAdata::RuffeSLRH92','BACI: estudo de caso com limitações. Longitudinal: preservar unidades. O catálogo de planejamento não promete um motor de análise para cada desenho.'),
 ('TR007 TR008 TR009 TR010 TR011','Planejamento experimental','datasets::PlantGrowth; datasets::npk; agridat::cochran.latin; salvelino_formalina_remocao; MASS::oats','Conferir aleatorização, unidades, blocos e níveis; parcelas subdivididas exigem estratos de erro próprios.'),
 ('TR012 TR013 TR014 TR015 TR016','Dimensionamento amostral','FSAdata::WhiteGrunt2; FSAdata::Cabezon; morfometria_barbo','Bases piloto; valores de efeito/precisão escolhidos não são dados observados adicionais.'),
 ('TR017 TR018 TR019 TR020','Seleção da amostra','walleye_erie; pinguins; doubs_ambiente; FSAdata::BGHRsample + BGHRfish','Sorteio didático não transforma os dados disponíveis em censo representativo. Conglomerados: preservar hierarquia.'),
 ('TR021 TR022 TR023 TR024 TR025','Importação e organização','treino_coletas; treino_desembarque; tilapia_sensorial; enobieta_baci','Dois treinos sintéticos autorizados, sempre identificados. Não arrumar previamente a estrutura que constitui o exercício.'),
 ('TR026 TR027 TR028 TR029 TR030 TR031 TR032 TR033 TR034 TR035 TR036 TR037 TR038 TR039','Tratamento e bases derivadas','captura_pescada_amarela_bruta; treino_coletas; treino_desembarque; pinguins; biometria_caranguejos; FSAdata::Cabezon','Conferir cada defeito/coluna no exercício. Linhas iguais não provam duplicata; não inserir defeitos novos sem aprovação.'),
 ('TR040 TR041 TR042 TR043 TR044 TR045 TR046 TR047 TR048 TR049 TR050 TR051','Exploração e gráficos','abalone_adultos; pinguins; morfometria_barbo; brine_carbonatos; cpue_tubarao','Escolher colunas e grupos; verificar pressupostos no exemplo final, não apenas o formato.'),
 ('TR052 TR053 TR054 TR055','Frequências e proporções','camaroes_sexo; lagostas_kelp_sexo; FSAdata::Cabezon','Frequências dependem de todas as linhas; preservar contagens e denominadores.'),
 ('TR056','McNemar','MASS::bacteria','44 pares completos semanas 0/11; cálculo já conferido no motor. Preservar ID e perdas.'),
 ('TR057 TR058 TR059','Testes t','biometria_caranguejos; idades_savel_repetibilidade; FSAdata::WhiteGrunt2','Pareamento depende da unidade original; conferir independência e estrutura na seleção final.'),
 ('TR060 TR061 TR062 TR063 TR064','ANOVA, subamostras, repetidas e ANCOVA','isoproteica_bagre; whitefish_salinidade; datasets::CO2; salvelino_formalina_remocao; crabs_morfometria','Whitefish não é pareamento individual; tanques e plantas precisam permanecer no desenho. A aprovação não resolve modelo inadequado.'),
 ('TR065 TR066','Variâncias','datasets::InsectSprays; biometria_caranguejos','Examinar distribuição e sensibilidade à não normalidade; não escolher por p-valor desejado.'),
 ('TR067 TR068 TR069 TR070','Não paramétricos','darter_ontario; idades_savel_repetibilidade; datasets::InsectSprays; tilapia_sensorial','Tilápia: confirmar âncoras/sentido da escala; atributos separados e pós-testes pareados.'),
 ('TR071 TR072 TR073','Correlação e regressões linear/logística','morfometria_barbo; abalone_adultos; FSAdata::Cabezon','Resposta binária válida e covariáveis; não tratar transformações como novas medições.'),
 ('TR074 TR075','Contagens Poisson/NB','datasets::InsectSprays; MASS::ships; MASS::Insurance; MASS::quine','Ships: exposição; NB específica teve limite de iterações. InsectSprays conferido; quine é alternativa a validar no motor.'),
 ('TR076 TR077 TR078 TR079 TR080 TR081 TR082','Curvas e crescimento','MASS::wtloss; bagley_lwr_central_america; snapper_hg; datasets::Orange; datasets::cars; datasets::Puromycin','Puromycin: log como aproximação, saturação como uso adicional. Orange: árvores repetidas. Fontes/unidades/pressupostos ainda sujeitos à curadoria.'),
 ('TR083 TR084 TR085','Séries temporais','astsa::rec; cpue_tubarao; exportacao_pargo','Recrutamento mensal para sazonalidade; não atribuir decomposição mensal a séries anuais curtas.'),
 ('TR086 TR087','PCA e agrupamento','brine_carbonatos; peixes_nutricao_portugal; enobieta_baci; doubs_ambiente','Distinguir matriz de abundância/presença da ambiental; padronização e distância adequadas.'),
 ('TR088 TR089 TR090','Mapas disponíveis','datasets::quakes; aquicultura_br','Quakes: exemplo real não pesqueiro. Doubs xy é coordenada local, não latitude/longitude; não converter por suposição.'),
 ('TR091 TR092 TR093','Conceitos, simuladores e comunicação','pinguins; morfometria_barbo; datasets::PlantGrowth','TLC pode sortear da base existente. Simuladores geram valores em execução; não adicionar bases sintéticas. Exportação é fluxo, não exige nova base.'),
]
covered={t for ids,*_ in coverage_groups for t in ids.split()}
assert covered=={f'TR{i:03}' for i in range(1,94)}

pending=[
 ['NOAA ERSSTv6','Enriquecimento opcional','Não obtido; recorte espacial/temporal a definir. Não bloqueia séries temporais nem os três mapas atuais. Decisão N003 anterior preservada no arquivo original.'],
 ['Mingau com pó de peixe','Busca encerrada por substituição','Tilápia sensorial cobre o uso pretendido. Não há dados brutos locais do mingau e ele não entra na contagem.'],
 ['8 objetos extras FESta','Fora da proposta','GAMstd, MGMSstd e SVstd documentados como simulados; outros cinco com fonte insuficientemente clara. Cópias ficam no estágio, sem admissão.'],
 ['readxl datasets.xlsx','Arquivo auxiliar opcional','Não contado como conjunto novo; usar as planilhas existentes para importação.'],
 ['BACI de maior replicação','Melhoria opcional','Enobieta é um caso real útil, mas não tem múltiplas datas por fase nem barragens independentes replicadas.'],
 ['12 entradas TR094-TR105','Recursos em preparação','Fora do fechamento das 93 entradas disponíveis; dados para funcionalidades futuras serão reavaliados quando houver implementação.'],
 ['Procedência, licenças e pressupostos','Curadoria obrigatória','Pronto na fonte não significa aprovado para incorporar. Alertas automáticos de fonte precisam leitura humana.'],
]
data={'data':'10/10/2026','unidade_contagem':'Uma entrada por objeto/conjunto proposto; estudos de literatura multiarquivo agrupados. Não é número de estudos independentes.',
 'total_registros':len(rows),'nucleo_proposto':len(core),'reservas':36,'apoio':3,'sinteticos_extras':6,
 'bytes_inventario':total,'bytes_nucleo':corebytes,'decisoes_anteriores':dict(Counter(x['decisao'] for x in decisions.values())),
 'arquivos_manifesto_conferidos':len(manifest),'divergencias_manifesto':changes,'hashes_entradas':hashes,'conjuntos':rows,'cobertura':coverage_groups,'pendencias':pending}
(ROOT/'balanco_levantamento_fechado.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')

wb=Workbook(); ws=wb.active; ws.title='Leia primeiro'
overview=[
 ['Balanço do levantamento - fotografia de 10/10/2026','Conteúdo'],
 ['Núcleo proposto para a Fase 2',f'60 entradas: 34 atuais + 5 fontes de literatura/base aberta + 21 fontes R. É proposta de seleção, não corte aplicado.'],
 ['Inventário completo','105 entradas: 43 atuais + 5 literatura/bases abertas + 57 R (inclui Puromycin e dois extras).'],
 ['Reserva e apoio','36 alternativas; 3 objetos de apoio; 6 sintéticos extras existentes aguardando decisão.'],
 ['Unidade de contagem','Objeto/conjunto proposto; whitefish, mesocosmos e Enobieta agrupam suas tabelas. Há fontes sobrepostas; não são 105 estudos independentes.'],
 ['Como decidir','Edite somente Decisão Fase 2 e Avaliação final do autor na aba Inventário. Núcleo 60 é uma visão de consulta, não uma segunda fila de decisões.'],
 ['Decisões anteriores','Copiadas por nome, com comentários, para referência. O XLSX original permanece intacto e é a fonte das aprovações anteriores. A seleção Fase 2 tem escopo novo e não cancela aprovações.'],
 ['Exemplo de preenchimento','Na linha de tilapia_sensorial, escolha Aprovar com ajustes e escreva: confirmar as âncoras da escala no artigo antes da adaptação. É somente um exemplo; não foi preenchido como decisão.'],
 ['Pronto','Nos candidatos R significa disponível na origem, com adequação e adaptação ainda sujeitas à revisão.'],
 ['Tamanho do inventário',f'{total:,} bytes = {total/1e6:.3f} MB. RDA atuais em disco, RDA candidatos em xz, três XLSX completos no formato recebido. Estimativa mista conservadora, não tamanho final do pacote.'],
 ['Tamanho do núcleo',f'{corebytes:,} bytes = {corebytes/1e6:.3f} MB, pelo mesmo critério, sem compressão adicional das três planilhas.'],
 ['Não incluído no tamanho','Código/documentação do pacote, arquivos fonte/tarballs, dependências, atividade PDF e fontes opcionais não obtidas. MB decimal = 1.000.000 bytes.'],
 ['Cobertura','Há candidatos para as 93 entradas disponíveis do catálogo; inclui tratamentos, planejamento, conceitos e exportação, não 93 testes diferentes. Homologação completa continua pendente.'],
 ['Integridade',f'{len(manifest)} arquivos do manifesto original conferidos: zero divergências. Nenhum dado atual alterado, removido ou incorporado.'],
 ['Base dos números','Fotografia calculada dos inventários JSON, arquivos locais e planilhas recebidas; totais desta aba são valores estáticos de auditoria. Não é planilha de cálculo dinâmica.'],
 ['Fase 2','Avaliar seleção, origem/licença, desenho, nomes, níveis e ajustes; executar somente alterações aprovadas, com documentação das adaptações.'],
]
for r in overview: ws.append(r)
headers=['nome do conjunto','origem','tipo de análise atendida','situação','avaliação resumida do que vai mudar','Papel no balanço','Decisão Fase 2','Avaliação final do autor','Decisão anterior (referência)','Comentário anterior (referência)','ID do balanço','ID anterior','Dimensão','Natureza','Bytes medidos/estimados','Fonte']
keys=['nome','origem','analise','situacao','avaliacao','papel','decisao_fase2','avaliacao_autor','decisao_anterior','comentario_anterior','id','id_anterior','dimensao','natureza','bytes','fonte']
for name,items in [('Núcleo 60',core),('Inventário',rows)]:
    s=wb.create_sheet(name); s.append(headers)
    for r in items: s.append([r[k] for k in keys])
    s.freeze_panes='F2'; s.auto_filter.ref=s.dimensions
    if name=='Inventário':
        dv=DataValidation(type='list',formula1='"Pendente,Aprovar proposta,Aprovar com ajustes,Recusar,Adiar"')
        dv.error='Escolha uma opção da lista.'; dv.showErrorMessage=True
        s.add_data_validation(dv); dv.add(f'G2:G{len(items)+1}')
        for row in s.iter_rows(min_row=2):
            for cell in [row[6],row[7]]: cell.fill=PatternFill('solid',fgColor='E6F1FC'); cell.font=Font(name='Arial',size=10,color='1464A0')
s=wb.create_sheet('Cobertura');s.append(['IDs do catálogo','Família','Candidatos no núcleo','Condição/limite'])
for r in coverage_groups:s.append(r)
s=wb.create_sheet('Pendências');s.append(['Item','Classificação','Avaliação'])
for r in pending:s.append(r)
for s in wb:
    s.sheet_view.showGridLines=False
    s.freeze_panes=s.freeze_panes or 'A2'
    for cell in s[1]: cell.fill=PatternFill('solid',fgColor='0F3B5F');cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
    s.row_dimensions[1].height=31
    for row in s.iter_rows(min_row=2):
        for cell in row:
            if cell.font.name!='Arial':cell.font=Font(name='Arial',size=10,color='253748')
            cell.alignment=Alignment(vertical='top',wrap_text=True)
        s.row_dimensions[row[0].row].height=95 if s.title in ['Inventário','Núcleo 60'] else 64
    widths=([34,28,43,19,85,29,24,65,25,60,17,16,29,32,24,68] if s.title in ['Inventário','Núcleo 60'] else [38,100] if s.title=='Leia primeiro' else [36,35,75,90] if s.title=='Cobertura' else [36,36,110])
    from openpyxl.utils import get_column_letter
    for i,width in enumerate(widths,1):s.column_dimensions[get_column_letter(i)].width=width
    if s.title!='Leia primeiro':s.auto_filter.ref=s.dimensions
dest=ROOT/'balanco_fase2_60_conjuntos.xlsx'
if dest.exists():raise RuntimeError('O balanço já existe; não sobrescrever decisões.')
wb.save(dest)
check=load_workbook(dest,read_only=True,data_only=False)
assert check['Núcleo 60'].max_row==61 and check['Inventário'].max_row==106
assert not any(c.data_type=='e' for s in check for row in s for c in row)
assert all(sha(Path(p))==h for p,h in hashes.items())

md=f'''# Fechamento do levantamento e proposta de núcleo para a Fase 2

Fotografia de 10/10/2026. Busca ampla encerrada; nenhuma incorporação ou corte aplicado.

## Número proposto: 60 conjuntos

| Grupo | Entradas |
|---|---:|
| Conjuntos atuais indicados para a curadoria analítica | 34 |
| Fontes novas de literatura/bases abertas | 5 |
| Candidatos de pacotes R para o núcleo | 21 |
| **Núcleo proposto para a Fase 2** | **60** |

Os 34 atuais correspondem a 43 objetos existentes menos seis sintéticos extras e três objetos de apoio. Essa é uma classificação da proposta, não remoção. Os dois treinos sintéticos autorizados estão entre os 34. Os cinco novos estudos/fontes são whitefish, mesocosmos, Artemia, tilápia sensorial e Enobieta. Os 21 candidatos R estão listados na planilha, com Puromycin incluído uma única vez, como conjunto do R fornecido em XLSX.

## Inventário completo preservado: 105 entradas

43 atuais + 5 literatura/bases abertas + 57 fontes R = 105. Sua distribuição de prioridade é: 60 no núcleo proposto, 36 em reserva, três de apoio e seis sintéticos extras para decisão. Os arquivos dos oito outros objetos FESta, o exemplo readxl e fontes ainda não obtidas não entram nessas 105 entradas.

Uma entrada é um objeto/conjunto proposto, não um artigo ou estudo independente. Whitefish reúne quatro tabelas; mesocosmos reúne seis; Enobieta reúne duas. BGHRsample e BGHRfish são dois objetos relacionados. Existem sobreposições como pinguins/penguins e doubs_ambiente/doubs: ficam visíveis para escolher formatos e usos na curadoria, sem alegar novos estudos independentes.

## Cobertura e qualidade

Há candidatos para as 93 entradas disponíveis do catálogo da TRILHA, agrupadas em 20 famílias na aba Cobertura. O catálogo inclui planejamento, preparo, exploração, análises, conceitos e exportação; não representa 93 testes diferentes. As 12 entradas em preparação não entram nessa conclusão. Isso certifica a identificação de candidatos e estruturas, não pressupostos nem homologação completa dos fluxos.

- **McNemar:** MASS::bacteria, com 44 pares completos semanas 0/11 e perdas documentadas. Cálculo conferido anteriormente no motor da TRILHA.
- **Friedman sensorial:** tilápia, 122 avaliadores com três tratamentos e dois atributos. Confirmar o significado e o sentido da escala antes de interpretar aceitação.
- **BACI:** Enobieta cobre o estudo de caso e a atividade prospectiva de planejamento. Uma intervenção e uma campanha por fase limitam inferência; as subamostras não são barragens independentes.
- **Logarítmico:** Puromycin cobre a aproximação empírica no intervalo observado; Michaelis-Menten é o uso principal de saturação. O ajuste logarítmico não tem assíntota.
- **Contagens e exposição:** InsectSprays, ships, Insurance e quine fornecem alternativas. Captura FESta em toneladas não é contagem Poisson. O ajuste NB específico de ships teve limite de iterações, registrado na curadoria.
- **Mapas e séries:** quakes, aquicultura_br e rec cobrem os usos atuais. NOAA e base pesqueira georreferenciada adicional são enriquecimentos opcionais.
- **Fontes e licenças:** bases com origem incompleta, resultados agregados e alertas de possível simulação/digitalização continuam sinalizadas. Alerta automático não equivale a confirmação; aprovação de nomes não resolve procedência.

## Tamanho

| Escopo | Bytes | MB decimal |
|---|---:|---:|
| Núcleo proposto de 60 | {corebytes:,} | {corebytes/1e6:.3f} |
| Inventário completo de 105 | {total:,} | {total/1e6:.3f} |

Critério: arquivos atuais `.rda` no tamanho real; candidatos previamente medidos como `.rda` em xz; tilápia, Enobieta e Puromycin contados como XLSX completos, incluindo documentação. Nenhum arquivo foi recomprimido agora. É uma estimativa mista conservadora para os dados reunidos, não o tamanho instalado ou compactado de um pacote final. Não inclui scripts, documentação do pacote, dependências, tarballs de origem, PDF da atividade ou fontes opcionais futuras. O inventário de dados medido fica {1_000_000-total:,} bytes abaixo de 1 MB.

## Revisão e integridade

Abra `balanco_fase2_60_conjuntos.xlsx`. A aba Núcleo 60 apresenta a proposta; a aba Inventário preserva também reservas, apoio e sintéticos extras, com campos azuis de Decisão Fase 2/Avaliação final do autor. Decisões anteriores e comentários foram copiados por nome somente como referência; a planilha original não foi sobrescrita. A nova seleção não cancela aprovações anteriores nem autoriza mudanças ainda não aprovadas.

{len(manifest)} arquivos do manifesto original conferidos, zero divergências. As cinco planilhas de entrada foram preservadas por SHA256. A Fase 2 começa pela revisão deste núcleo e dos limites científicos; qualquer adaptação aplicada depois deve informar fonte, licença e alterações de nomes, categorias, unidades, colunas e linhas.
'''
(ROOT/'balanco_levantamento_fechado.md').write_text(md,encoding='utf-8')
print(json.dumps({k:v for k,v in data.items() if k not in ['conjuntos','cobertura','pendencias','hashes_entradas']},ensure_ascii=False,indent=2))
print(dest)
